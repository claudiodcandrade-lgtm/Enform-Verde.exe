import sys, json, math, tempfile, re, zipfile, threading, queue, traceback, base64, io, os, requests
from pathlib import Path
import tkinter as tk
import webbrowser
from tkinter import ttk, filedialog, messagebox
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageTk
from sar_pipeline import discover_sar, process_real_sar, automatic_pipeline, MODEL_REGISTRY, model_registry_rows, scientific_calibration_report, cdse_access_token
from lband_preprocess import preprocess_lband

APP_VERSION="3.24.15-PROFESSIONAL"
ORANGE="#EF9B06"; FOREST="#0B3D2E"; GREEN="#155D43"; PALE="#F4F6F5"; TEXT="#34413E"

# Fontes implementadas no motor. Valores-proxy são sempre rotulados como MODELADOS.
SOURCES={
 "protocol":"Higa et al. (2014), Embrapa Florestas, Documentos 266",
 "soil":"PronaSolos/Embrapa Solos, estoque de carbono orgânico 90 m, perfis 0–30/60/100/200 cm + diagnóstico de incerteza",
 "deadwood":"Freitas et al. (2021), Embrapa Amazônia Ocidental — necromassa lenhosa",
 "litter":"Embrapa Amazônia Oriental — estudos de serapilheira; proxy só para triagem",
}
ROOT_RATIO=0.26; ROOT_LOW=0.18; ROOT_HIGH=0.30
CARBON_FRACTION=0.47

def app_resource(name):
    base=Path(getattr(sys,"_MEIPASS",Path(sys.executable).parent)) if getattr(sys,"frozen",False) else Path(__file__).parent
    p=base/name
    if p.exists():return p
    # onedir builds may keep data beside the executable rather than inside _internal
    q=Path(sys.executable).parent/name if getattr(sys,"frozen",False) else p
    return q

OFFLINE_BRAZIL_BOUNDS=(-75.0,-35.0,-33.0,6.0)  # west,south,east,north

def offline_brazil_preview(gdf,extent_factor=1.36,out_size=(1000,600)):
    """Render AOI over the packaged NASA Blue Marble Brazil mosaic without network access."""
    path=app_resource("offline_brazil_basemap.jpg")
    if not path.exists():raise FileNotFoundError("offline_brazil_basemap.jpg ausente")
    g=gdf.to_crs(4326).copy(); minx,miny,maxx,maxy=map(float,g.total_bounds)
    dx=max(maxx-minx,1e-8); dy=max(maxy-miny,1e-8); f=max(0.30,min(8.0,float(extent_factor)))
    cx=(minx+maxx)/2; cy=(miny+maxy)/2
    bx=[cx-dx*f/2,cy-dy*f/2,cx+dx*f/2,cy+dy*f/2]
    W,S,E,N=OFFLINE_BRAZIL_BOUNDS
    bx=[max(W,bx[0]),max(S,bx[1]),min(E,bx[2]),min(N,bx[3])]
    if bx[0]>=bx[2] or bx[1]>=bx[3]:raise ValueError("AOI fora da cobertura do mosaico offline do Brasil")
    im=Image.open(path).convert("RGB"); iw,ih=im.size
    def srcxy(lon,lat):
        x=(lon-W)/(E-W)*iw; y=(N-lat)/(N-S)*ih
        return x,y
    x0,y0=srcxy(bx[0],bx[3]); x1,y1=srcxy(bx[2],bx[1])
    crop=im.crop((max(0,int(x0)),max(0,int(y0)),min(iw,int(math.ceil(x1))),min(ih,int(math.ceil(y1)))))
    ow,oh=map(int,out_size); crop=crop.resize((ow,oh),Image.Resampling.LANCZOS)
    from PIL import ImageDraw
    d=ImageDraw.Draw(crop,"RGBA")
    for geom in g.geometry:
        geoms=list(geom.geoms) if geom.geom_type=="MultiPolygon" else ([geom] if geom.geom_type=="Polygon" else [])
        for poly in geoms:
            pts=[]
            for coord in poly.exterior.coords:
                lon,lat=float(coord[0]),float(coord[1])
                x=(lon-bx[0])/max(bx[2]-bx[0],1e-12)*ow
                y=(bx[3]-lat)/max(bx[3]-bx[1],1e-12)*oh
                pts.append((x,y))
            if len(pts)>=4:
                d.polygon(pts,fill=(255,138,0,42),outline=(255,138,0,255))
                d.line(pts,width=max(3,round(ow/400)),fill=(255,138,0,255),joint="curve")
    return crop,bx


def agb_mexiana(dbh_cm):
    return 0.1184*np.power(np.asarray(dbh_cm,dtype=float),2.53)

def load_inventory(path):
    df=pd.read_excel(path,sheet_name="Biomassa árvores vivas",usecols="A:E")
    df.columns=["plot","tree","dbh_m","height_m","species"]; df["plot"]=df["plot"].ffill()
    df=df[df["tree"].notna() & df["dbh_m"].notna()].copy()
    df["dbh_cm"]=pd.to_numeric(df["dbh_m"],errors="coerce")*100
    df["height_m"]=pd.to_numeric(df["height_m"],errors="coerce")
    return df[df["dbh_cm"].between(10,400)].copy()

def summarize_inventory(df,plot_area_m2=250):
    d=df.copy(); d["agb_kg"]=agb_mexiana(d["dbh_cm"])
    p=d.groupby("plot",as_index=False)["agb_kg"].sum()
    x=(p["agb_kg"]/1000*(10000/plot_area_m2)).to_numpy(float)
    if len(x)<2: raise ValueError("São necessárias ao menos 2 parcelas válidas.")
    mean=float(np.mean(x)); se=float(np.std(x,ddof=1)/math.sqrt(len(x)))
    return len(x),mean,max(0,mean-1.96*se),mean+1.96*se

def spatial_libs():
    try:
        import geopandas as gpd
        return gpd
    except Exception as e: raise RuntimeError("Módulo espacial não disponível: "+str(e))

def read_vector(path):
    gpd=spatial_libs(); p=Path(path)
    if p.suffix.lower()==".kmz":
        d=Path(tempfile.mkdtemp(prefix="enform_kmz_"))
        with zipfile.ZipFile(path) as z:
            ks=[i for i in z.infolist() if not i.is_dir() and i.filename.lower().endswith(".kml")]
            if not ks: raise ValueError("KMZ sem arquivo KML interno.")
            # Only the first KML document is needed. Never extract arbitrary archive
            # paths (a crafted KMZ could otherwise overwrite files outside the temp dir).
            info=ks[0]
            if info.file_size>50*1024*1024: raise ValueError("O KML interno do KMZ excede 50 MiB.")
            target=d/"document.kml"
            with z.open(info) as src, target.open("wb") as dst:
                while True:
                    chunk=src.read(1024*1024)
                    if not chunk: break
                    dst.write(chunk)
        gdf=gpd.read_file(target,driver="KML")
    elif p.suffix.lower()==".kml":
        gdf=gpd.read_file(path,driver="KML")
    else:gdf=gpd.read_file(path)
    if gdf.empty: raise ValueError("O vetor não contém feições.")
    if gdf.crs is None: raise ValueError("O vetor não informa CRS; não é seguro assumir coordenadas.")
    gdf=gdf[gdf.geometry.notna() & ~gdf.geometry.is_empty].copy()
    if gdf.empty: raise ValueError("Não há geometria válida.")
    return gdf

def geom_metrics(gdf):
    g=gdf.to_crs(4326); union=g.geometry.union_all()
    c=union.centroid; lon,lat=float(c.x),float(c.y)
    zone=max(1,min(60,int((lon+180)//6)+1)); epsg=(32700 if lat<0 else 32600)+zone
    gm=gdf.to_crs(epsg); area=float(gm.geometry.area.sum()/10000)
    b=union.bounds
    return {"area_ha":area,"centroid":[lon,lat],"bbox":[float(x) for x in b],"utm_epsg":epsg}

def _sicar_session():
    import requests, ssl
    from requests.adapters import HTTPAdapter
    class SicarTLSAdapter(HTTPAdapter):
        def init_poolmanager(self,*args,**kwargs):
            from urllib3.poolmanager import PoolManager
            ctx=ssl.create_default_context()
            try: ctx.set_ciphers("DEFAULT@SECLEVEL=1")
            except Exception: pass
            kwargs["ssl_context"]=ctx
            return super().init_poolmanager(*args,**kwargs)
    sess=requests.Session()
    sess.mount("https://geoserver.car.gov.br/",SicarTLSAdapter())
    sess.headers.update({"User-Agent":f"Enform-Verde/{APP_VERSION}","Accept":"application/json"})
    return sess

def resolve_car(car):
    """Resolve CAR against the official public SICAR WFS, with protocol fallbacks and diagnostics."""
    import geopandas as gpd, requests
    raw=str(car or "").strip().upper().replace("–","-").replace("—","-")
    code="".join(raw.split())
    compact="".join(ch for ch in code if ch.isalnum())
    if len(compact)==41:
        code=compact[:2]+"-"+compact[2:9]+"-"+compact[9:]
    m=re.match(r"^([A-Z]{2})-([0-9]{7})-([A-F0-9]{32})$",code)
    if not m: raise ValueError("Código CAR inválido/incompleto. Use o código integral no padrão UF-7 dígitos-32 caracteres.")
    uf=m.group(1).lower(); layer=f"sicar:sicar_imoveis_{uf}"
    urls=["https://geoserver.car.gov.br/geoserver/sicar/ows","https://geoserver.car.gov.br/geoserver/sicar/wfs"]
    attempts=[]; sess=_sicar_session()
    for url in urls:
      for version,key in [("1.0.0","typeName"),("1.1.0","typeName"),("2.0.0","typeNames")]:
        params={"service":"WFS","version":version,"request":"GetFeature",key:layer,
                "outputFormat":"application/json","srsName":"EPSG:4326","cql_filter":f"cod_imovel='{code}'"}
        try:
            try:r=sess.get(url,params=params,timeout=(8,35))
            except requests.exceptions.SSLError:r=sess.get(url,params=params,timeout=(8,35),verify=False)
            attempts.append(f"{version}:{r.status_code}")
            if not r.ok:continue
            js=r.json(); feats=js.get("features") or []
            if not feats:continue
            exact=[x for x in feats if str((x.get("properties") or {}).get("cod_imovel","")).upper()==code]
            feats=exact or feats
            if len(feats)>1:raise RuntimeError("SICAR retornou registros múltiplos para o mesmo código.")
            props=feats[0].get("properties") or {}
            if not feats[0].get("geometry"):raise RuntimeError("CAR localizado, porém sem geometria pública.")
            tmp=Path(tempfile.gettempdir())/"enform_car_sicar.geojson"
            tmp.write_text(json.dumps({"type":"FeatureCollection","features":feats}),encoding="utf-8")
            gdf=gpd.read_file(tmp); gdf=gdf.set_crs(4326) if gdf.crs is None else gdf.to_crs(4326)
            gdf.attrs.update({"sicar_cod_imovel":props.get("cod_imovel",code),"sicar_status":props.get("status_imovel") or props.get("condicao"),
              "sicar_area_declarada_ha":props.get("area"),"sicar_municipio":props.get("municipio"),"sicar_lookup":url+" WFS "+version})
            return gdf
        except RuntimeError:raise
        except Exception as e:attempts.append(version+":"+type(e).__name__)
    raise LookupError("CAR válido no formato, mas não localizado no WFS público oficial nesta consulta. Tentativas: "+", ".join(attempts)+". O programa não concluirá que o CAR é inválido por indisponibilidade do serviço.")

def resolve_ccir_sigef(code):
    """Resolve código INCRA/SNCR do CCIR em parcela georreferenciada certificada no SIGEF."""
    import requests, geopandas as gpd
    digits=re.sub(r"\D","",code or "")
    if len(digits)!=13: raise ValueError("Informe o Código do Imóvel Rural de 13 dígitos constante do CCIR.")
    # Consulta pública de parcelas do SIGEF pelo código do imóvel (SNCR/INCRA).
    search="https://sigef.incra.gov.br/api/parcelas/consulta/"
    sess=requests.Session(); sess.headers.update({"User-Agent":f"Enform-Verde/{APP_VERSION}","Accept":"application/json"})
    candidates=[]
    for key in ("codigo_imovel","codigo_imovel_incra","sncr"):
        try:
            r=sess.get(search,params={key:digits,"format":"json"},timeout=(10,45))
            if not r.ok: continue
            js=r.json()
            if isinstance(js,dict):
                candidates=js.get("results") or js.get("features") or js.get("parcelas") or []
            elif isinstance(js,list): candidates=js
            if candidates: break
        except Exception: continue
    if not candidates:
        raise LookupError("Nenhuma parcela certificada no SIGEF foi localizada automaticamente para o código INCRA/SNCR do CCIR. O CCIR é cadastral e só possui geometria quando há correspondência georreferenciada no SIGEF/SNCI.")
    item=candidates[0]
    if isinstance(item,dict) and item.get("geometry"):
        tmp=Path(tempfile.gettempdir())/"enform_sigef.geojson"
        tmp.write_text(json.dumps({"type":"FeatureCollection","features":[item]}),encoding="utf-8")
        gdf=gpd.read_file(tmp)
        return gdf.to_crs("EPSG:4326") if gdf.crs else gdf.set_crs("EPSG:4326")
    url=item.get("geojson") or item.get("geometry_url") or item.get("download") if isinstance(item,dict) else None
    if url:
        r=sess.get(url,timeout=(10,60)); r.raise_for_status()
        tmp=Path(tempfile.gettempdir())/"enform_sigef.geojson"; tmp.write_bytes(r.content)
        gdf=gpd.read_file(tmp)
        return gdf.to_crs("EPSG:4326") if gdf.crs else gdf.set_crs("EPSG:4326")
    raise RuntimeError("A parcela foi localizada no SIGEF, mas a consulta pública não forneceu geometria em formato utilizável automaticamente.")

IBGE_VEGE_2026_URL="https://geoftp.ibge.gov.br/informacoes_ambientais/vegetacao/vetores/escala_250_mil/versao_2026/vege_area.zip"
IBGE_BIOMAS_2025_URL="https://geoftp.ibge.gov.br/informacoes_ambientais/estudos_ambientais/biomas/vetores/2025_Biomas-e-Sistema-Costeiro-Marinho-do-Brasil-1-250000_shp.zip"

def _ibge_cache():
    import os
    root=Path(os.environ.get("LOCALAPPDATA",Path.home()))/"EnformVerde"/"IBGE"
    root.mkdir(parents=True,exist_ok=True); return root

def _download_zip(url,folder,tag):
    """Cache the official ZIP without extracting the national dataset into RAM/disk."""
    import requests
    folder.mkdir(parents=True,exist_ok=True)
    zpath=folder/(tag+".zip")
    if zpath.exists() and zpath.stat().st_size>1024*1024:return zpath
    tmp=zpath.with_suffix(".part")
    with requests.get(url,stream=True,timeout=(15,300),headers={"User-Agent":f"Enform-Verde/{APP_VERSION}"}) as r:
        r.raise_for_status()
        with open(tmp,"wb") as out:
            for chunk in r.iter_content(4*1024*1024):
                if chunk:out.write(chunk)
    tmp.replace(zpath)
    return zpath

def _read_zip_bbox(zpath,bbox4326,preferred=("area",)):
    """Read only polygons intersecting the property bbox directly from ZIP (GDAL /vsizip)."""
    import pyogrio
    from pyproj import Transformer
    vsi="zip://"+Path(zpath).as_posix()
    layers=pyogrio.list_layers(vsi)
    candidates=[]
    for row in layers:
        name=str(row[0]); geom=str(row[1] or "")
        if "polygon" in geom.lower():
            score=sum(10 for x in preferred if x.lower() in name.lower())+(5 if "brasil" in name.lower() else 0)
            candidates.append((score,name))
    if not candidates:raise RuntimeError("Pacote IBGE sem camada poligonal reconhecida.")
    layer=max(candidates)[1]
    info=pyogrio.read_info(vsi,layer=layer)
    crs=info.get("crs")
    bb=tuple(float(x) for x in bbox4326)
    if crs and str(crs).upper() not in ("EPSG:4326","EPSG:4674"):
        tr=Transformer.from_crs("EPSG:4326",crs,always_xy=True)
        x1,y1=tr.transform(bb[0],bb[1]); x2,y2=tr.transform(bb[2],bb[3]); bb=(min(x1,x2),min(y1,y2),max(x1,x2),max(y1,y2))
    g=pyogrio.read_dataframe(vsi,layer=layer,bbox=bb,use_arrow=False)
    if g.empty:raise RuntimeError("A camada IBGE não retornou polígonos no envelope da propriedade.")
    return g

def _field(cols,candidates):
    low={str(c).lower():c for c in cols}
    for x in candidates:
        if x.lower() in low:return low[x.lower()]
    for c in cols:
        lc=str(c).lower()
        if any(x.lower() in lc for x in candidates):return c
    return None

def _shares(project,theme,fields):
    import geopandas as gpd
    p=project.to_crs("EPSG:5880"); t=theme.to_crs("EPSG:5880")
    try:
        import shapely
        p["geometry"]=p.geometry.map(lambda g:shapely.make_valid(g) if g is not None and not g.is_valid else g)
        t["geometry"]=t.geometry.map(lambda g:shapely.make_valid(g) if g is not None and not g.is_valid else g)
    except Exception: pass
    inter=gpd.overlay(t,p[["geometry"]],how="intersection",keep_geom_type=False)
    inter=inter[inter.geometry.notna() & ~inter.geometry.is_empty].copy()
    if inter.empty:return []
    inter["_ha"]=inter.geometry.area/10000
    total=float(inter["_ha"].sum())
    if total<=0:return []
    out=[]
    for field in fields:
        if field and field in inter.columns:
            g=inter.groupby(field,dropna=False)["_ha"].sum().sort_values(ascending=False)
            out.append((field,[(str(k),float(v),float(v/total*100)) for k,v in g.items() if v>0]))
    return out

def _shares_with_code(project,theme,name_field,code_field):
    """Area-weighted IBGE region names retaining their official phytoregion code."""
    import geopandas as gpd
    p=project.to_crs("EPSG:5880"); t=theme.to_crs("EPSG:5880")
    try:
        import shapely
        p["geometry"]=p.geometry.map(lambda g:shapely.make_valid(g) if g is not None and not g.is_valid else g)
        t["geometry"]=t.geometry.map(lambda g:shapely.make_valid(g) if g is not None and not g.is_valid else g)
    except Exception: pass
    columns=["geometry",name_field]+([code_field] if code_field else [])
    inter=gpd.overlay(t[columns],p[["geometry"]],how="intersection",keep_geom_type=False)
    inter=inter[inter.geometry.notna() & ~inter.geometry.is_empty].copy()
    if inter.empty:return []
    inter["_ha"]=inter.geometry.area/10000; total=float(inter["_ha"].sum())
    if total<=0:return []
    keys=[name_field]+([code_field] if code_field else [])
    out=inter.groupby(keys,dropna=False)["_ha"].sum().sort_values(ascending=False)
    return [{"name":str(key[0] if isinstance(key,tuple) else key),
             "code":(str(key[1]) if code_field and isinstance(key,tuple) and key[1] is not None else None),
             "area_ha":float(area),"percent":float(area/total*100)} for key,area in out.items() if area>0]

def primary_ibge_physiognomy(diagnosis):
    """Return the dominant current IBGE phytogeographic region (legend_1).

    legend_2 describes the predominant vegetation cover class and is useful as
    a secondary descriptor, but it is not a substitute for the phytogeographic
    region requested for model-domain selection.
    """
    for group in (diagnosis or {}).get("vegetacao",[]):
        if str(group.get("campo","")).casefold()=="legenda_1":
            classes=group.get("classes") or []
            if classes:return str(classes[0][0])
    return ""

def diagnose_ibge(project):
    """IBGE diagnosis with bbox-only reads; national 355 MB vegetation ZIP is never expanded/read wholesale."""
    root=_ibge_cache(); veg=root/"vegetacao_2026"; bio=root/"biomas_2025"
    bbox=tuple(project.to_crs("EPSG:4326").total_bounds)
    bz=_download_zip(IBGE_BIOMAS_2025_URL,bio,"biomas_2025")
    bg=_read_zip_bbox(bz,bbox,("bioma","area"))
    bfield=_field(bg.columns,["Bioma","Nome_Bioma","nm_bioma","bioma_1"])
    bs=_shares(project,bg,[bfield])
    if not bs or not bs[0][1]:raise RuntimeError("O polígono não interceptou a camada oficial de Biomas do IBGE.")
    result={"bioma_field":bfield,"biomas":bs[0][1],"vegetacao_fields":[],"vegetacao":[],
            "fonte_bioma":"IBGE — Biomas do Brasil, revisão 2025, 1:250.000",
            "fonte_vegetacao":"IBGE — Vegetação/Regiões Fitoecológicas, versão 2026, 1:250.000"}
    try:
        vz=_download_zip(IBGE_VEGE_2026_URL,veg,"vege_area_2026")
        vg=_read_zip_bbox(vz,bbox,("vege_area","area","brasil"))
        # In the official BDIA 1:250,000 schema, legenda_1 is the name of the
        # phytogeographic region; legenda_2 is the predominant cover class.
        # Do not silently substitute the latter for the former.
        l1=_field(vg.columns,["legenda_1"])
        l2=_field(vg.columns,["legenda_2"])
        vs=_shares(project,vg,[l1,l2])
        result["vegetacao_fields"]=[x[0] for x in vs]
        result["vegetacao"]=[{"campo":x[0],"classes":x[1]} for x in vs]
        fito_code=_field(vg.columns,["cd_fito"])
        result["regioes_fitoecologicas"]=_shares_with_code(project,vg,l1,fito_code) if l1 else []
        if not any(x[0]=="legenda_1" and x[1] for x in vs):
            result["vegetacao_error"]="A camada foi lida, mas não retornou Região Fitoecológica IBGE em legenda_1 para a AOI. A análise não deve tratar a cobertura de legenda_2 como substituta."
    except Exception as e:
        result["vegetacao_error"]=str(e)
    return result

def zonal_soil(gdf,raster_path):
    import rasterio
    from rasterio.mask import mask
    gg=gdf.to_crs(rasterio.open(raster_path).crs)
    shapes=[x.__geo_interface__ for x in gg.geometry]
    with rasterio.open(raster_path) as src:
        a,_=mask(src,shapes,crop=True,filled=False)
        vals=a[0].compressed().astype(float)
        if src.nodata is not None: vals=vals[vals!=src.nodata]
    if len(vals)==0: raise ValueError("O raster de solo não possui pixels válidos na área.")
    return float(np.mean(vals)),float(np.std(vals)),len(vals)

def pronasolos_soc_profiles(gdf,max_points=36):
    """PronaSolos/Embrapa Solos 90 m SOC stocks (Mg C/ha), cumulative profiles.
    Uses only the verified official MapServer. Each depth is queried independently."""
    import requests
    from shapely.geometry import Point
    gg=gdf.to_crs(3857); geom=gg.geometry.union_all()
    minx,miny,maxx,maxy=geom.bounds
    # systematic polygon sampling plus representative point; avoid dependence on bbox corners
    n=max(4,int(math.ceil(math.sqrt(max_points))))
    xs=np.linspace(minx,maxx,n); ys=np.linspace(miny,maxy,n)
    pts=[Point(float(x),float(y)) for y in ys for x in xs if geom.covers(Point(float(x),float(y)))]
    rp=geom.representative_point()
    pts=[rp]+pts
    # de-duplicate and cap
    uniq=[]; seen=set()
    for p in pts:
        k=(round(p.x,2),round(p.y,2))
        if k not in seen: seen.add(k); uniq.append(p)
    pts=uniq[:max_points]
    service="https://geoportal.sgb.gov.br/server/rest/services/pronasolos/estoque_carbono_90m/MapServer/identify"
    layer_ids=[2,3,4,5,6,7]
    names=["0–5","5–15","15–30","30–60","60–100","100–200"]
    vals={k:[] for k in layer_ids}; errors={k:[] for k in layer_ids}
    sess=requests.Session()
    extent=f"{minx},{miny},{maxx},{maxy}"
    def numeric_value(item):
        at=item.get("attributes") or {}
        candidates=[item.get("value")]
        # ArcGIS raster identify varies by server/version/language.
        for key,val in at.items():
            kl=str(key).lower()
            if ("pixel" in kl and "value" in kl) or "stretched" in kl or kl in ("value","valor","pixel_value"):
                candidates.append(val)
        for raw in candidates:
            if raw is None: continue
            try:
                txt=str(raw).strip().replace(" ","").replace(",",".")
                v=float(txt)
                if np.isfinite(v) and 0 <= v < 2000: return v
            except Exception: pass
        return None
    for p in pts:
        # Query each layer independently: one missing/deep layer can never invalidate 0-30.
        for lid in layer_ids:
            params={"f":"json","geometry":f"{p.x},{p.y}","geometryType":"esriGeometryPoint","sr":"3857",
                    "layers":f"all:{lid}","tolerance":"5","mapExtent":extent,
                    "imageDisplay":"1600,1600,96","returnGeometry":"false"}
            try:
                r=sess.get(service,params=params,timeout=(5,20)); r.raise_for_status(); js=r.json()
                if js.get("error"): raise RuntimeError(str(js["error"]))
                found=False
                for item in js.get("results",[]):
                    if int(item.get("layerId",-1))!=lid: continue
                    v=numeric_value(item)
                    if v is not None: vals[lid].append(v); found=True; break
                if not found: errors[lid].append("sem valor numérico")
            except Exception as e: errors[lid].append(str(e)[:160])
    targets={"0–30 cm":3,"0–60 cm":4,"0–100 cm":5,"0–200 cm":6}; out={}
    for label,count in targets.items():
        need=layer_ids[:count]
        if any(not vals[k] for k in need): continue
        # Use depth means independently; this avoids pairing samples incorrectly when one depth has a local NoData.
        means=[float(np.mean(vals[k])) for k in need]
        stock=float(sum(means))
        # Conservative descriptive spatial propagation; NOT prediction uncertainty/CI95.
        sds=[float(np.std(vals[k],ddof=1)) if len(vals[k])>1 else 0.0 for k in need]
        spatial_sd=float(math.sqrt(sum(x*x for x in sds)))
        out[label]={"tc_ha":stock,"spatial_sd_tc_ha":spatial_sd,
          "n_samples":min(len(vals[k]) for k in need),"layers":names[:count],
          "layer_means_tc_ha":dict(zip(names[:count],means)),
          "source":"PronaSolos/Embrapa Solos — estoque de carbono orgânico 90 m",
          "uncertainty_kind":"DP espacial propagado das amostras do mapa; não é IC95% nem erro de predição"}
    if "0–30 cm" not in out:
        detail=[]
        for i,k in enumerate(layer_ids[:3]):
            detail.append(f"{names[i]} cm: n={len(vals[k])}; "+(("; ".join(errors[k][:2])) if errors[k] else "sem retorno"))
        raise RuntimeError("PronaSolos oficial consultado, mas 0–30 cm não pôde ser composto. "+" | ".join(detail))
    out["_diagnostic"]={"samples_requested":len(pts),"valid_by_layer":{names[i]:len(vals[k]) for i,k in enumerate(layer_ids)},
                        "service":service}
    return out


def self_test():
    assert abs(float(agb_mexiana(10))-0.1184*10**2.53)<1e-8
    assert ROOT_LOW<ROOT_RATIO<ROOT_HIGH
    assert abs(CARBON_FRACTION-0.47)<1e-9
    if sys.stdout is not None:print("ENFORM_VERDE_SELF_TEST_OK")

def acceptance_test(kmz_path):
    """Run the real Tapajós SAR acceptance path from the compiled executable."""
    gdf=read_vector(kmz_path)
    if gdf.empty:raise RuntimeError("AOI de aceitação sem feições.")
    metrics=geom_metrics(gdf)
    result=automatic_pipeline(gdf,"Amazônia","Floresta Ombrófila Densa das Terras Baixas",
                              cache=str(Path(tempfile.gettempdir())/"enform_verde_acceptance_sar"))
    audit=result.get("audit") or {}
    processed=list(audit.get("processed_without_agb") or [])
    pixel_count=sum(int(s.get("n",0)) for item in processed for s in item.get("stats",[]))
    pixel_count+=sum(int(s.get("n",0)) for s in result.get("stats",[]))
    if pixel_count<=0:raise RuntimeError("Nenhum pixel SAR foi processado dentro da AOI Tapajós.")
    agb=result.get("agb_mg_ha")
    if agb is None:
        lit=result.get("literature_reference") or {}
        if not lit.get("available") or not lit.get("sar_processed"):
            raise RuntimeError("Pixels SAR foram processados, mas a referência regional Tapajós não foi selecionada corretamente.")
    elif not str(result.get("data_origin","")).startswith("SAR"):
        raise RuntimeError("Origem AGB incompatível: resultado numérico sem proveniência SAR.")
    summary={"status":result.get("status"),"area_ha":metrics["area_ha"],"pixel_count":pixel_count,
             "data_origin":result.get("data_origin"),"scene_count":len(audit.get("alos_palsar_public",{}).get("scene_ids",[]))+len(audit.get("sentinel1_public",{}).get("scene_ids",[])),
             "fallback_agb_mg_ha":(result.get("literature_reference") or {}).get("agb_mg_ha")}
    (Path(tempfile.gettempdir())/"enform_verde_acceptance_result.json").write_text(json.dumps(summary,ensure_ascii=True),encoding="utf-8")
    return summary

def ui_smoke_test():
    root=App(); root.after(1200,root.destroy); root.mainloop()

def sentinel2_preview(gdf,out_h=700,extent_factor=1.36):
    """Recent low-cloud Sentinel-2 L2A RGB preview and AOI pixel coordinates; no user API key."""
    import rasterio
    from rasterio.warp import transform_bounds, transform
    from rasterio.windows import from_bounds
    g=gdf.to_crs(4326).copy(); geom=g.geometry.union_all().__geo_interface__
    body={"collections":["sentinel-2-l2a"],"intersects":geom,"limit":30,
          "sortby":[{"field":"properties.datetime","direction":"desc"}]}
    rr=requests.post("https://planetarycomputer.microsoft.com/api/stac/v1/search",json=body,timeout=(10,60)); rr.raise_for_status()
    items=rr.json().get("features",[]); cand=[]
    for it in items:
        visual=(it.get("assets") or {}).get("visual")
        if not visual or not visual.get("href"):continue
        props=it.get("properties") or {}
        cloud=float(props.get("eo:cloud_cover",100.0) if props.get("eo:cloud_cover") is not None else 100.0)
        dt=str(props.get("datetime") or "")
        cand.append((cloud,dt,it,visual["href"]))
    if not cand:raise RuntimeError("Nenhuma cena Sentinel-2 L2A RGB encontrada para a AOI.")
    low=[x for x in cand if x[0]<=20.0]
    chosen=sorted(low,key=lambda x:x[1],reverse=True)[0] if low else sorted(cand,key=lambda x:(x[0],x[1]))[0]
    cloud,dt,it,unsigned=chosen
    sg=requests.get("https://planetarycomputer.microsoft.com/api/sas/v1/sign",params={"href":unsigned},timeout=(10,45)); sg.raise_for_status(); href=sg.json()["href"]
    with rasterio.Env(GDAL_HTTP_MULTIRANGE="YES",GDAL_HTTP_MERGE_CONSECUTIVE_RANGES="YES"):
        with rasterio.open(href) as src:
            b=list(map(float,g.total_bounds)); bx=transform_bounds("EPSG:4326",src.crs,*b,densify_pts=21)
            dx=max(bx[2]-bx[0],1.0); dy=max(bx[3]-bx[1],1.0)
            factor=max(0.30,min(8.0,float(extent_factor))); cx=(bx[0]+bx[2])/2; cy=(bx[1]+bx[3])/2
            wb=(cx-dx*factor/2,cy-dy*factor/2,cx+dx*factor/2,cy+dy*factor/2)
            win=from_bounds(*wb,transform=src.transform).round_offsets().round_lengths()
            out_w=max(700,min(1200,round(out_h*max(float(win.width),1.0)/max(float(win.height),1.0))))
            arr=src.read(indexes=[1,2,3],window=win,out_shape=(3,out_h,out_w),
                         resampling=rasterio.enums.Resampling.bilinear,boundless=True,fill_value=0)
            rgb=np.moveaxis(arr,0,2)
            if rgb.dtype!=np.uint8:
                valid=rgb[np.isfinite(rgb)&(rgb>0)]
                hi=float(np.percentile(valid,99)) if valid.size else 1.0
                rgb=np.clip(rgb/max(hi,1e-9)*255,0,255).astype(np.uint8)
            view=Image.fromarray(rgb,"RGB"); wt=src.window_transform(win); polygons=[]
            for geom0 in g.geometry:
                geoms=list(geom0.geoms) if geom0.geom_type=="MultiPolygon" else ([geom0] if geom0.geom_type=="Polygon" else [])
                for poly in geoms:
                    coords=list(poly.exterior.coords)
                    lon=[p[0] for p in coords]; lat=[p[1] for p in coords]
                    xx,yy=transform("EPSG:4326",src.crs,lon,lat); pts=[]
                    for x,y in zip(xx,yy):
                        col,row=(~wt)*(x,y)
                        pts.extend([float(col)*out_w/max(float(win.width),1.0),float(row)*out_h/max(float(win.height),1.0)])
                    if len(pts)>=6:polygons.append(pts)
    return {"image":view,"polygons":polygons,"scene_id":it.get("id"),"datetime":dt,"cloud_cover":cloud,
            "provider":"Sentinel-2 L2A / Microsoft Planetary Computer","asset":"visual"}

class App(tk.Tk):
    def __init__(self):
        super().__init__(); self.title("Enform Verde"); screen_w=self.winfo_screenwidth(); screen_h=self.winfo_screenheight(); win_w=max(1100,min(1713,screen_w-48)); win_h=max(620,min(918,screen_h-88)); self.geometry(f"{win_w}x{win_h}"); self.minsize(min(1024,win_w),min(600,win_h))
        self.inv=None; self.gdf=None; self.soil_raster=None; self.project={"version":APP_VERSION}; self.active_source=None; self.active_input_id=None; self._analysis_running=False; self._analysis_queue=queue.Queue(); self._ibge_queue=queue.Queue(); self._ibge_generation=0; self._ibge_pending=False; self._pending_execute=False; self._map_generation=0; self._map_refresh_job=None; self._map_queue=queue.Queue()
        self._style(); self._ui(); self.bind("<Return>",self.execute)
    def _style(self):
        s=ttk.Style(self)
        try:s.theme_use("vista")
        except:pass
        s.configure(".",font=("Segoe UI",10),foreground=TEXT)
        s.configure("Title.TLabel",font=("Segoe UI",20,"bold"),foreground=ORANGE)
        s.configure("H.TLabel",font=("Segoe UI",12,"bold"),foreground=FOREST)
        s.configure("Run.TButton",font=("Segoe UI",10,"bold"),padding=10)
        s.configure("TButton",padding=7)
    def _ui(self):
        # Professional dashboard shell based on the approved Enform Verde reference.
        root=ttk.Frame(self); root.pack(fill="both",expand=True)
        base=app_resource(".").resolve()

        # Approved high-resolution Enform mask; render once, preserve aspect ratio, add no text overlays.
        header_h=210
        header=tk.Canvas(root,height=header_h,bg="#063D26",highlightthickness=0); header.pack(fill="x",side="top")
        visual=base/"enform_header.png"
        self.header_source=Image.open(visual).convert("RGB") if visual.exists() else None
        def render_header(event=None):
            w=max(1,int(event.width if event else header.winfo_width() or 1400)); h=max(1,int(event.height if event else header_h))
            header.delete("all")
            if self.header_source is None:
                return
            # Aspect ratio is preserved; the whole approved mask stays visible at every window size.
            scale=min(w/self.header_source.width,h/self.header_source.height)
            iw=max(1,round(self.header_source.width*scale)); ih=max(1,round(self.header_source.height*scale))
            im=self.header_source.resize((iw,ih),Image.Resampling.LANCZOS)
            self.header_photo=ImageTk.PhotoImage(im)
            header.create_image((w-iw)//2,(h-ih)//2,image=self.header_photo,anchor="nw")
        header.bind("<Configure>",render_header); render_header()

        # Global action bar: EXECUTAR ANÁLISE must remain visible regardless of selected section.
        action=ttk.Frame(root,padding=(305,12,32,10)); action.pack(fill="x",side="top")
        ttk.Label(action,text="Análise de carbono",style="Title.TLabel").pack(side="left")
        ttk.Button(action,text="Salvar relatório",command=self.save_report).pack(side="right",padx=(8,0))
        ttk.Button(action,text="Exportar Excel",command=self.export_excel).pack(side="right",padx=(8,0))
        self.global_execute_btn=ttk.Button(action,text="EXECUTAR ANÁLISE",command=self.execute,style="Run.TButton")
        self.global_execute_btn.pack(side="right",padx=(8,0))

        body=ttk.Frame(root); body.pack(fill="both",expand=True)
        nav=tk.Frame(body,width=288,bg="#F6F8F8",highlightbackground="#D8E0E0",highlightthickness=1)
        nav.pack(side="left",fill="y"); nav.pack_propagate(False)
        main=ttk.Frame(body,padding=(16,8,32,6)); main.pack(side="left",fill="both",expand=True)

        self.nb=ttk.Notebook(main); self.nb.pack(fill="both",expand=True)
        self.tabs=[]
        for n in ["Propriedade (CAR)","Mapas e limites","Carregar produtos SAR / AGB","Carbono Total","Base Científica e Modelos"]:
            f=ttk.Frame(self.nb,padding=14); self.nb.add(f,text=n); self.tabs.append(f)

        nav_items=[
          ("⌂   Início",0),("▣   Propriedade (CAR)",0),("◫   Mapas e limites",1),
          ("◈   Fitofisionomia (IBGE)",1),("◆   Carregar produtos\n     SAR / AGB",2),
          ("◉   Biomassa Aérea (AGB)",3),("♨   Carbono no Solo",3),
          ("●   Carbono Total",3),("▤   Resultados e Relatório",3),
          ("⚙   Configurações",2),("▥   Base Científica e\n     Modelos",4),("ⓘ   Sobre",4)]
        self.nav_buttons=[]
        for label,idx in nav_items:
            btn=tk.Button(nav,text=label,anchor="w",justify="left",relief="flat",bd=0,
                          bg="#F6F8F8",fg="#233B49",activebackground="#E5F0EB",
                          font=("Segoe UI",10),padx=20,pady=11,
                          command=lambda i=idx:self.nb.select(i))
            btn.pack(fill="x",pady=0); self.nav_buttons.append((btn,idx))
        def mark_tab(event=None):
            cur=self.nb.index(self.nb.select())
            for btn,idx in self.nav_buttons:
                if idx==cur:
                    btn.configure(bg="#08733F",fg="white",font=("Segoe UI",10,"bold"))
                else:
                    btn.configure(bg="#F6F8F8",fg="#233B49",font=("Segoe UI",10))
        self.nb.bind("<<NotebookTabChanged>>",mark_tab)

        self._project(); self._spatial(); self._remote(); self._results(); self._sources()
        mark_tab()
        footer=ttk.Frame(root,padding=(14,4)); footer.pack(fill="x",side="bottom")
        ttk.Label(footer,text="v"+APP_VERSION).pack(side="left")
        self.status=tk.StringVar(value="Pronto. Informe o CAR ou carregue o vetor da propriedade.")
        ttk.Label(footer,textvariable=self.status,anchor="center").pack(side="left",fill="x",expand=True)
        ttk.Label(footer,text="Sistema de Estimativa de Estoques de Carbono em Vegetação Nativa").pack(side="right")

    def _project(self):
        f=self.tabs[0]; ttk.Label(f,text="Abrir análise",style="H.TLabel").grid(row=0,column=0,columnspan=4,sticky="w")
        self.name=tk.StringVar(value="Projeto Enform Verde"); self.car=tk.StringVar(); self.ccir=tk.StringVar()
        self.biome=tk.StringVar(value=""); self.phys=tk.StringVar(value="")
        fields=[("Projeto",self.name),("CAR / SICAR",self.car),("CCIR — código INCRA/SNCR (13 dígitos)",self.ccir)]
        for i,(lab,var) in enumerate(fields,1):
            ttk.Label(f,text=lab).grid(row=i,column=0,sticky="w",pady=8)
            ttk.Entry(f,textvariable=var,width=62).grid(row=i,column=1,sticky="ew",padx=10)
        ttk.Label(f,text="A consulta CAR/SICAR ou CCIR/SIGEF ocorre automaticamente ao pressionar EXECUTAR ANÁLISE.",foreground="#52645E",wraplength=900).grid(row=4,column=0,columnspan=3,sticky="w",pady=(12,4))
        ttk.Button(f,text="CARREGAR ARQUIVO VETORIAL",command=self.pick_vector,style="Run.TButton").grid(row=5,column=1,sticky="w",pady=10,padx=10)
        ttk.Label(f,text="KML • KMZ • SHP • GeoJSON • GPKG",foreground="#666").grid(row=5,column=2,sticky="w")
        self.vector_status=tk.StringVar(value="Nenhum arquivo vetorial carregado.")
        ttk.Label(f,textvariable=self.vector_status,foreground="#155D43",wraplength=900).grid(row=6,column=1,columnspan=3,sticky="w",padx=10,pady=(2,8))
        f.columnconfigure(1,weight=1)

    def _spatial(self):
        host=self.tabs[1]
        tab_canvas=tk.Canvas(host,bg="#F4F6F5",highlightthickness=0)
        tab_v=ttk.Scrollbar(host,orient="vertical",command=tab_canvas.yview)
        tab_canvas.configure(yscrollcommand=tab_v.set)
        tab_canvas.pack(side="left",fill="both",expand=True); tab_v.pack(side="right",fill="y")
        f=ttk.Frame(tab_canvas,padding=(0,0,10,10))
        win=tab_canvas.create_window((0,0),window=f,anchor="nw")
        def sync_scroll(event=None):
            tab_canvas.configure(scrollregion=tab_canvas.bbox("all"))
            tab_canvas.itemconfigure(win,width=max(700,tab_canvas.winfo_width()))
        f.bind("<Configure>",sync_scroll); tab_canvas.bind("<Configure>",sync_scroll)

        ttk.Label(f,text="Perímetro, diagnóstico e solo",style="H.TLabel").pack(anchor="w")
        row=ttk.Frame(f); row.pack(fill="x",pady=8)
        ttk.Button(row,text="Buscar COS 0–30 cm — Embrapa",command=self.auto_soil).pack(side="left")
        ttk.Button(row,text="Carregar GeoTIFF de COS",command=self.pick_soil).pack(side="left",padx=8)
        ttk.Button(row,text="VISUALIZAR SATÉLITE + POLÍGONO",command=self.show_google_map,style="Run.TButton").pack(side="left",padx=8)
        ttk.Button(row,text="−",width=3,command=self.map_zoom_out).pack(side="left",padx=(12,2))
        ttk.Button(row,text="+",width=3,command=self.map_zoom_in).pack(side="left",padx=2)
        ttk.Button(row,text="AJUSTAR AOI",command=self.map_zoom_fit).pack(side="left",padx=(2,8))
        keyrow=ttk.Frame(f); keyrow.pack(fill="x",pady=(0,6))
        ttk.Label(keyrow,text="Chave Google Maps Static API:").pack(side="left")
        self.google_maps_key=tk.StringVar(value=os.environ.get("GOOGLE_MAPS_API_KEY",""))
        ttk.Entry(keyrow,textvariable=self.google_maps_key,show="•",width=48).pack(side="left",padx=8)
        ttk.Label(keyrow,text="Google opcional; fallback: Sentinel-2 → Esri Imagery → Esri Street → NASA Blue Marble offline",foreground="#666").pack(side="left")

        ibgebox=ttk.LabelFrame(f,text="Classificação oficial IBGE",padding=(10,7)); ibgebox.pack(fill="x",pady=(0,6))
        self.ibge_biome_display=tk.StringVar(value="Aguardando perímetro.")
        self.ibge_phys_display=tk.StringVar(value="Aguardando classificação legenda_1.")
        ttk.Label(ibgebox,text="Bioma IBGE:",font=("Segoe UI",9,"bold")).grid(row=0,column=0,sticky="nw")
        ttk.Label(ibgebox,textvariable=self.ibge_biome_display,wraplength=840).grid(row=0,column=1,sticky="w",padx=(8,0))
        ttk.Label(ibgebox,text="Fitofisionomia IBGE (legenda_1):",font=("Segoe UI",9,"bold")).grid(row=1,column=0,sticky="nw",pady=(4,0))
        ttk.Label(ibgebox,textvariable=self.ibge_phys_display,wraplength=840,foreground="#155D43").grid(row=1,column=1,sticky="w",padx=(8,0),pady=(4,0))
        ibgebox.columnconfigure(1,weight=1)

        map_frame=ttk.Frame(f); map_frame.pack(fill="both",expand=True,pady=(2,6))
        self.map_canvas=tk.Canvas(map_frame,height=420,bg="#DDE4E1",highlightthickness=1,highlightbackground="#B8C5C0",
                                  xscrollincrement=20,yscrollincrement=20)
        self.map_hscroll=ttk.Scrollbar(map_frame,orient="horizontal",command=self.map_canvas.xview)
        self.map_vscroll=ttk.Scrollbar(map_frame,orient="vertical",command=self.map_canvas.yview)
        self.map_canvas.configure(xscrollcommand=self.map_hscroll.set,yscrollcommand=self.map_vscroll.set)
        self.map_canvas.grid(row=0,column=0,sticky="nsew"); self.map_vscroll.grid(row=0,column=1,sticky="ns"); self.map_hscroll.grid(row=1,column=0,sticky="ew")
        map_frame.rowconfigure(0,weight=1); map_frame.columnconfigure(0,weight=1)
        self._map_redraw_job=None; self._satellite_map_visible=False; self._last_map_size=None; self._map_extent_factor=1.36
        self.map_canvas.bind("<Configure>",self._on_map_resize)
        self.map_canvas.create_text(20,20,anchor="nw",text="Carregue CAR, CCIR ou vetor e visualize o satélite com o polígono.",fill="#455")
        self.map_canvas.configure(scrollregion=self.map_canvas.bbox("all"))

        spatial_box=ttk.Frame(f); spatial_box.pack(fill="x",pady=4)
        self.spatial_text=tk.Text(spatial_box,height=7,wrap="word",yscrollcommand=lambda *a:spatial_scroll.set(*a))
        spatial_scroll=ttk.Scrollbar(spatial_box,orient="vertical",command=self.spatial_text.yview)
        self.spatial_text.pack(side="left",fill="both",expand=True); spatial_scroll.pack(side="right",fill="y")
        self._set(self.spatial_text,"Nenhum perímetro carregado. Use CAR, CCIR/SIGEF ou arquivo vetorial na tela de abertura.")

    def map_zoom_in(self):
        if self.gdf is None:return
        self._map_extent_factor=max(0.30,self._map_extent_factor/1.5)
        self.show_google_map(refresh=True,prefer_public=True)

    def map_zoom_out(self):
        if self.gdf is None:return
        self._map_extent_factor=min(8.0,self._map_extent_factor*1.5)
        self.show_google_map(refresh=True,prefer_public=True)

    def map_zoom_fit(self):
        if self.gdf is None:return
        self._map_extent_factor=1.36
        self.show_google_map(refresh=True,prefer_public=True)

    def show_google_map(self,refresh=False,prefer_public=False):
        """Satellite visualization: Google when keyed, then Sentinel-2, then Esri, then offline vector."""
        if self.gdf is None:return messagebox.showwarning("Mapa","Carregue/resolva o polígono primeiro.")
        self._map_generation+=1; generation=self._map_generation
        g=self.gdf.to_crs(4326).copy()
        viewport_w=max(500,self.map_canvas.winfo_width()); viewport_h=max(280,self.map_canvas.winfo_height())
        display_scale=max(1.0,min(4.0,1.36/max(float(self._map_extent_factor),0.30)))
        w=max(500,round(viewport_w*display_scale)); h=max(280,round(viewport_h*display_scale))
        ratio=w/max(h,1)
        if ratio>=1:
            req_w=1200; req_h=max(320,min(1200,round(1200/ratio)))
        else:
            req_h=1200; req_w=max(320,min(1200,round(1200*ratio)))
        key=self.google_maps_key.get().strip(); extent_factor=max(0.30,min(8.0,float(self._map_extent_factor)))
        if not refresh:self.status.set("Carregando imagem de satélite e perímetro...")
        self.update_idletasks()

        def rings():
            out=[]
            span=max(float(g.total_bounds[2]-g.total_bounds[0]),float(g.total_bounds[3]-g.total_bounds[1]),1e-7)
            tol=max(span/1600.0,1e-7)
            for geom in g.geometry:
                geoms=list(geom.geoms) if geom.geom_type=="MultiPolygon" else ([geom] if geom.geom_type=="Polygon" else [])
                for poly in geoms:
                    coords=list(poly.simplify(tol,preserve_topology=True).exterior.coords)
                    if len(coords)>220:
                        step=max(1,math.ceil(len(coords)/220)); coords=coords[::step]
                        if coords[-1]!=coords[0]:coords.append(coords[0])
                    if len(coords)>=4:out.append(coords)
            return out

        def google_image(rs):
            if not key:return None
            params=[("size",f"{min(req_w,640)}x{min(req_h,640)}"),("scale","2"),("maptype","satellite"),("format","png"),("key",key)]
            for coords in rs:
                pts="|".join(f"{lat:.6f},{lon:.6f}" for lon,lat in coords)
                params.append(("path","color:0xff8a00ff|weight:4|fillcolor:0xff8a0033|"+pts))
            rr=requests.get("https://maps.googleapis.com/maps/api/staticmap",params=params,timeout=(8,35))
            rr.raise_for_status()
            if "image" not in (rr.headers.get("Content-Type") or "").lower():
                raise RuntimeError("resposta Google não é imagem")
            return Image.open(io.BytesIO(rr.content)).convert("RGB").resize((w,h),Image.Resampling.LANCZOS)

        def esri_export(rs,service="World_Imagery",label="Esri"):

            minx,miny,maxx,maxy=map(float,g.total_bounds)
            dx=max(maxx-minx,1e-8); dy=max(maxy-miny,1e-8); cx=(minx+maxx)/2; cy=(miny+maxy)/2
            bx=(cx-dx*extent_factor/2,cy-dy*extent_factor/2,cx+dx*extent_factor/2,cy+dy*extent_factor/2)
            params={"bbox":",".join(f"{v:.8f}" for v in bx),"bboxSR":"4326","imageSR":"4326",
                    "size":f"{req_w},{req_h}","format":"jpg","f":"image","dpi":"96"}
            rr=requests.get(f"https://server.arcgisonline.com/ArcGIS/rest/services/{service}/MapServer/export",params=params,timeout=(8,35))
            rr.raise_for_status()
            if "image" not in (rr.headers.get("Content-Type") or "").lower():
                raise RuntimeError("resposta Esri não é imagem")
            im=Image.open(io.BytesIO(rr.content)).convert("RGB")
            from PIL import ImageDraw
            d=ImageDraw.Draw(im,"RGBA"); iw,ih=im.size
            def px(lon,lat):
                return ((lon-bx[0])/max(bx[2]-bx[0],1e-12)*iw,
                        (bx[3]-lat)/max(bx[3]-bx[1],1e-12)*ih)
            for coords in rs:
                pts=[px(float(lon),float(lat)) for lon,lat in coords]
                if len(pts)>=4:
                    d.polygon(pts,fill=(255,138,0,48),outline=(255,138,0,255))
                    d.line(pts,width=max(3,round(iw/400)),fill=(255,138,0,255),joint="curve")
            return im.resize((w,h),Image.Resampling.LANCZOS)

        def worker():
            errors=[]; rs=rings()
            if not rs:
                self._map_queue.put((generation,None,[],"Mapa",None,None,"geometria sem polígono utilizável"))
                return
            if key and not prefer_public:
                try:
                    view=google_image(rs)
                    self._map_queue.put((generation,view,[],"Google Maps Static API",None,"Google",None))
                    return
                except Exception as e:errors.append("Google: "+str(e))
            try:
                p=sentinel2_preview(g,extent_factor=extent_factor)
                copyright=f"Sentinel-2 L2A • {(p.get('datetime') or '')[:10]} • nuvens da cena {float(p.get('cloud_cover',0)):.1f}%"
                self._map_queue.put((generation,p["image"],p["polygons"],p["provider"],None,copyright,None))
                return
            except Exception as e:errors.append("Sentinel-2: "+str(e))
            try:
                view=esri_export(rs,"World_Imagery","Esri World Imagery")
                self._map_queue.put((generation,view,[],"Esri World Imagery",None,"Esri / World Imagery",None))
                return
            except Exception as e:errors.append("Esri Imagery: "+str(e))
            try:
                view=esri_export(rs,"World_Street_Map","Esri World Street Map")
                self._map_queue.put((generation,view,[],"Esri World Street Map",None,"Esri / World Street Map",None))
                return
            except Exception as e:errors.append("Esri Street: "+str(e))
            try:
                view,_=offline_brazil_preview(g,extent_factor,(w,h))
                self._map_queue.put((generation,view,[],"NASA Blue Marble offline",None,"NASA Blue Marble — fundo nacional offline",None))
                return
            except Exception as e:errors.append("Offline NASA: "+str(e))
            self._map_queue.put((generation,None,[],"Google/Sentinel-2/Esri/NASA",None,None," | ".join(errors)))

        threading.Thread(target=worker,name="EnformMap",daemon=True).start()
        self.after(80,self._poll_map_queue)

    def _poll_map_queue(self):
        try:item=self._map_queue.get_nowait()
        except queue.Empty:
            if self.winfo_exists():self.after(100,self._poll_map_queue)
            return
        generation,view,polygons,provider,zoom,copyright,error=item
        self._finish_satellite_map(generation,view,polygons,provider,zoom,copyright,error)

    def _finish_satellite_map(self,generation,view,polygons,provider,zoom,copyright,error):
        if generation!=self._map_generation or self.gdf is None:return
        if error:
            self._satellite_map_visible=False
            if not self._draw_offline_brazil_basemap("Mapa offline — NASA Blue Marble"):
                self._draw_aoi_outline("Perímetro carregado — visualização vetorial")
            self.status.set("Fontes online indisponíveis; fundo nacional offline mantido.")
            self._set(self.spatial_text,self.spatial_text.get("1.0","end").strip()+"\n\nMapa base indisponível: "+error)
            return
        self._satellite_map_visible=True; self._last_map_size=(view.width,view.height)
        self.google_map_photo=ImageTk.PhotoImage(view); self.map_canvas.delete("all"); self.map_canvas.create_image(0,0,image=self.google_map_photo,anchor="nw")
        for pts in polygons:self.map_canvas.create_polygon(*pts,fill="",outline="#FF8A00",width=3)
        self.map_canvas.create_rectangle(0,view.height-24,view.width,view.height,fill="white",outline=""); self.map_canvas.create_text(view.width-8,view.height-12,anchor="e",text=copyright,fill="#333",font=("Segoe UI",8))
        self.map_canvas.configure(scrollregion=(0,0,view.width,view.height))
        self.map_canvas.xview_moveto(max(0.0,min(1.0,(view.width-max(1,self.map_canvas.winfo_width()))/(2*max(1,view.width)))))
        self.map_canvas.yview_moveto(max(0.0,min(1.0,(view.height-max(1,self.map_canvas.winfo_height()))/(2*max(1,view.height)))))
        self.status.set(f"{provider} carregado com o perímetro. Zoom cartográfico {1.36/max(self._map_extent_factor,1e-9):.2f}×. Uso exclusivo para visualização.")

    def _on_map_resize(self,event=None):
        """Keep AOI and satellite base fitted after a real canvas resize."""
        if self.gdf is None:return
        size=(max(300,self.map_canvas.winfo_width()),max(220,self.map_canvas.winfo_height()))
        if size==self._last_map_size:return
        if self._satellite_map_visible:
            if self._map_refresh_job is not None:
                try:self.after_cancel(self._map_refresh_job)
                except Exception:pass
            self._map_refresh_job=self.after(450,self._refresh_satellite_after_resize)
            return
        if self._map_redraw_job is not None:
            try:self.after_cancel(self._map_redraw_job)
            except Exception:pass
        self._map_redraw_job=self.after(100,self._redraw_map_after_resize)

    def _refresh_satellite_after_resize(self):
        self._map_refresh_job=None
        if self.gdf is not None:
            self.show_google_map(refresh=True,prefer_public=True)

    def _redraw_map_after_resize(self):
        self._map_redraw_job=None
        if self.gdf is not None and not self._satellite_map_visible:
            self._draw_aoi_outline("AOI — visualização vetorial; base satélite indisponível")

    def _draw_offline_brazil_basemap(self,label="Mapa offline — NASA Blue Marble"):
        if self.gdf is None:return False
        try:
            viewport_w=max(500,self.map_canvas.winfo_width()); viewport_h=max(280,self.map_canvas.winfo_height())
            display_scale=max(1.0,min(4.0,1.36/max(float(self._map_extent_factor),0.30)))
            w=max(500,round(viewport_w*display_scale)); h=max(280,round(viewport_h*display_scale))
            view,_=offline_brazil_preview(self.gdf,self._map_extent_factor,(w,h))
            self._satellite_map_visible=True; self._last_map_size=(w,h)
            self.google_map_photo=ImageTk.PhotoImage(view); self.map_canvas.delete("all")
            self.map_canvas.create_image(0,0,image=self.google_map_photo,anchor="nw")
            self.map_canvas.create_rectangle(0,h-24,w,h,fill="white",outline="")
            self.map_canvas.create_text(w-8,h-12,anchor="e",text="NASA Blue Marble — fundo nacional offline",fill="#333",font=("Segoe UI",8))
            self.map_canvas.create_text(10,10,anchor="nw",text=label,fill="white",font=("Segoe UI",9,"bold"))
            self.map_canvas.configure(scrollregion=(0,0,w,h))
            return True
        except Exception:
            return False

    def _draw_aoi_outline(self,label="Pré-visualização da AOI"):
        """Render the AOI in a local metric projection and fit it to the live canvas."""
        if self.gdf is None:return
        self._satellite_map_visible=False
        g=self.gdf.to_crs(4326); union=g.geometry.union_all(); center=union.centroid
        from pyproj import CRS
        local=CRS.from_proj4(f"+proj=aeqd +lat_0={float(center.y)} +lon_0={float(center.x)} +datum=WGS84 +units=m +no_defs")
        projected=g.to_crs(local); minx,miny,maxx,maxy=map(float,projected.total_bounds)
        w=max(300,self.map_canvas.winfo_width()); h=max(220,self.map_canvas.winfo_height()); pad=max(30,min(70,int(min(w,h)*0.12)))
        dx=max(maxx-minx,1e-6); dy=max(maxy-miny,1e-6); scale=min((w-2*pad)/dx,(h-2*pad)/dy)
        draw_w=dx*scale; draw_h=dy*scale; ox=(w-draw_w)/2; oy=(h-draw_h)/2
        self._last_map_size=(w,h); self.map_canvas.delete("all"); self.map_canvas.configure(bg="#EAF0EC")
        # Offline map frame: always available and intentionally independent of tile/network services.
        for frac in (0.25,0.5,0.75):
            gx=pad+(w-2*pad)*frac; gy=pad+(h-2*pad)*frac
            self.map_canvas.create_line(gx,pad,gx,h-pad,fill="#C9D4CF",dash=(2,4))
            self.map_canvas.create_line(pad,gy,w-pad,gy,fill="#C9D4CF",dash=(2,4))
        self.map_canvas.create_text(w-14,14,anchor="ne",text="N",fill="#24382D",font=("Segoe UI",10,"bold"))
        self.map_canvas.create_line(w-20,48,w-20,24,fill="#24382D",width=2,arrow="first")
        def xy(x,y):return ox+(float(x)-minx)*scale,h-(oy+(float(y)-miny)*scale)
        for geom in projected.geometry:
            polys=list(geom.geoms) if geom.geom_type=="MultiPolygon" else ([geom] if geom.geom_type=="Polygon" else [])
            for poly in polys:
                pts=[]
                for coord in poly.exterior.coords:pts.extend(xy(coord[0],coord[1]))
                if len(pts)>=6:self.map_canvas.create_polygon(*pts,fill="#F8B44C",stipple="gray50",outline="#E87500",width=3)
        self.map_canvas.create_text(12,12,anchor="nw",text=label,fill="#24382D",font=("Segoe UI",10,"bold"))
        # Scale bar is calculated from the local projection; the AOI is never stretched to the panel.
        bar_m=max(1.0,draw_w*0.18/1000)*1000; bar_px=bar_m*scale; bx=max(18,w-pad-bar_px); by=h-24
        self.map_canvas.create_line(bx,by,bx+bar_px,by,fill="#24382D",width=3)
        self.map_canvas.create_line(bx,by-5,bx,by+4,fill="#24382D",width=2); self.map_canvas.create_line(bx+bar_px,by-5,bx+bar_px,by+4,fill="#24382D",width=2)
        self.map_canvas.create_text(bx+bar_px/2,by-7,text=f"{bar_m/1000:g} km",anchor="s",fill="#24382D",font=("Segoe UI",8,"bold"))
        b4326=g.total_bounds
        bbox_txt=f"{b4326[0]:.5f}, {b4326[1]:.5f}  →  {b4326[2]:.5f}, {b4326[3]:.5f}"
        self.map_canvas.create_text(12,h-12,anchor="sw",text=bbox_txt,fill="#4A5B54",font=("Segoe UI",8))

    def _remote(self):
        f=self.tabs[2]
        steps=ttk.Frame(f); steps.pack(fill="x",pady=(0,10))
        for j,t in enumerate(["1. Dados e Catálogos","2. Download e Pré-processamento","3. Modelagem e AGB","4. Resultados SAR"]):
            lab=tk.Label(steps,text=t,bg=("#08733F" if j==0 else "#EEF2F3"),fg=("white" if j==0 else "#233B49"),font=("Segoe UI",10,"bold"),padx=14,pady=9,bd=1,relief="solid")
            lab.pack(side="left",fill="x",expand=True,padx=(0,2))

        # Public routes remain automatic. Optional Earthdata Login unlocks protected NISAR/ASF scenes.
        self.edl_token=tk.StringVar(); self.edl_user=tk.StringVar(value=os.environ.get("EARTHDATA_USERNAME","")); self.edl_password=tk.StringVar()
        self.esa_token=tk.StringVar(); self.cdse_token=tk.StringVar(); self.cdse_client_id=tk.StringVar(); self.cdse_client_secret=tk.StringVar()
        edl=ttk.LabelFrame(f,text="Earthdata Login opcional — NISAR/ASF protegido",padding=8); edl.pack(fill="x",pady=(0,6))
        ttk.Label(edl,text="Usuário:").grid(row=0,column=0,sticky="w"); ttk.Entry(edl,textvariable=self.edl_user,width=28).grid(row=0,column=1,sticky="w",padx=(5,12))
        ttk.Label(edl,text="Senha:").grid(row=0,column=2,sticky="w"); ttk.Entry(edl,textvariable=self.edl_password,show="•",width=28).grid(row=0,column=3,sticky="w",padx=5)
        ttk.Label(edl,text="Não é armazenada. O programa também tenta EARTHDATA_* e _netrc/.netrc automaticamente.",foreground="#666").grid(row=1,column=0,columnspan=4,sticky="w",pady=(4,0))
        ttk.Label(edl,text="ESA MAAP Long Lasting Token (90 dias):").grid(row=2,column=0,sticky="w",pady=(5,0))
        ttk.Entry(edl,textvariable=self.esa_token,show="•",width=58).grid(row=2,column=1,columnspan=3,sticky="ew",padx=(5,5),pady=(5,0))
        ttk.Label(edl,text="Opcional; usado somente para baixar produtos BIOMASS protegidos e nunca gravado pelo programa.",foreground="#666").grid(row=3,column=0,columnspan=4,sticky="w",pady=(3,0))
        providers=ttk.LabelFrame(f,text="Fontes SAR automáticas — prioridade máxima",padding=12); providers.pack(fill="x",pady=(4,10))
        ttk.Label(providers,text="ESA BIOMASS — banda P",font=("Segoe UI",10,"bold")).grid(row=0,column=0,sticky="w")
        ttk.Label(providers,text="PRIORIDADE 1 — consulta automática das coleções BIOMASS L1/L2 e FP_AGB_L2B; processamento do produto disponível mais adequado.").grid(row=0,column=1,sticky="w",padx=8)
        ttk.Label(providers,text="ESA CCI Biomass v7",font=("Segoe UI",10,"bold")).grid(row=1,column=0,sticky="w",pady=5)
        ttk.Label(providers,text="AGB 100 m + incerteza por pixel — acesso automático como produto quantitativo SAR derivado.").grid(row=1,column=1,sticky="w",padx=8)
        ttk.Label(providers,text="Copernicus Sentinel-1",font=("Segoe UI",10,"bold")).grid(row=2,column=0,sticky="w")
        ttk.Label(providers,text="C-band VV/VH — consulta/processamento automático quando o serviço público permitir acesso direto.").grid(row=2,column=1,sticky="w",padx=8)
        ttk.Label(providers,text="NISAR / ALOS-PALSAR",font=("Segoe UI",10,"bold")).grid(row=3,column=0,sticky="w",pady=5)
        ttk.Label(providers,text="L-band — catálogo consultado automaticamente; produtos diretamente acessíveis são processados sem intervenção do usuário.").grid(row=3,column=1,sticky="w",padx=8)
        ttk.Label(providers,text="Arquivo local/licenciado",font=("Segoe UI",10,"bold")).grid(row=4,column=0,sticky="w")
        ttk.Label(providers,text="GeoTIFF/HDF5 SAR ou AGB continua disponível como rota adicional.").grid(row=4,column=1,sticky="w",padx=8)
        providers.columnconfigure(1,weight=1)

        row=ttk.Frame(f); row.pack(fill="x",pady=8)
        self.pipeline_btn=ttk.Button(row,text="EXECUTAR ANÁLISE SAR",command=self.execute,style="Run.TButton"); self.pipeline_btn.pack(side="left")
        ttk.Button(row,text="DESCOBRIR COBERTURA SAR",command=self.discover_sar_ui).pack(side="left",padx=8)
        ttk.Button(row,text="CARREGAR PRODUTOS SAR / AGB",command=self.pick_sar).pack(side="left")
        self.sar_paths=[]; self.sensor=tk.StringVar(value="Automático — SAR primeiro: P → L → X → C → CCI; literatura/modelagem somente após falha documentada")
        remote_box=ttk.Frame(f); remote_box.pack(fill="both",expand=True,pady=8)
        self.remote_text=tk.Text(remote_box,height=18,wrap="word",yscrollcommand=lambda *a:remote_scroll.set(*a))
        remote_scroll=ttk.Scrollbar(remote_box,orient="vertical",command=self.remote_text.yview)
        self.remote_text.pack(side="left",fill="both",expand=True); remote_scroll.pack(side="right",fill="y")
        self._set(self.remote_text,"A análise é automática e SAR-FIRST. Rotas públicas são tentadas sem credenciais; Earthdata Login é opcional para liberar NISAR/ASF protegido.\n\nHierarquia obrigatória:\n1. ESA BIOMASS FP_AGB_L2B (P-band, AGB + incerteza);\n2. modelos SAR L/X executáveis compatíveis com fitofisionomia e atributos disponíveis;\n3. ESA CCI Biomass L+C como série histórica;\n4. literatura somente como aferição/fallback quando nenhum produto SAR quantitativo puder ser processado.\n\nRegra: SAR é SEMPRE tentado primeiro. Se for impossível processá-lo, a trilha registra o motivo e só então usa literatura/modelagem compatível, identificada como secundária e com incerteza explícita.")

    def _test_cdse(self):
        try:
            cid=self.cdse_client_id.get().strip(); sec=self.cdse_client_secret.get().strip()
            if not cid or not sec:
                self.cdse_state.set("Informe Client ID e Client Secret.")
                return
            cdse_access_token(cid,sec)
            self.cdse_state.set("CDSE conectado — OAuth2 válido.")
        except Exception as e:
            self.cdse_state.set("Falha CDSE: "+str(e)[:120])

    def discover_sar_ui(self):
        if self.gdf is None:return messagebox.showwarning("SAR","Carregue/resolva o polígono primeiro.")
        try:
            self.status.set("Consultando catálogos SAR..."); self.update_idletasks()
            cov=discover_sar(self.gdf); self.project["sar_catalog"]=cov
            txt=["REGISTROS RETORNADOS PELOS CATÁLOGOS SAR","(descoberta ≠ download ≠ processamento ≠ uso na estimativa)"]
            for x in cov:txt.append(f"{x['provider']} — {x['dataset']} ({x['band']}): {x['count']} registro(s) retornado(s)"+((" — "+x["error"]) if x.get("error") else ""))
            txt += ["","Para uso quantitativo, o produto ainda precisa ser elegível, baixado, pré-processado e associado a modelo/produto AGB compatível."]
            self._set(self.remote_text,"\n".join(txt)); self.status.set("Descoberta SAR concluída.")
        except Exception as e:self.status.set("Falha na descoberta SAR."); messagebox.showerror("SAR",str(e))

    def pick_sar(self):
        ps=filedialog.askopenfilenames(title="Produtos SAR / AGB / incerteza",filetypes=[("GeoTIFF","*.tif *.tiff"),("Todos","*.*")])
        if not ps:return
        self.sar_paths=list(ps); self.project["sar_paths"]=list(ps)
        self._set(self.remote_text,"Produtos selecionados:\n"+"\n".join(self.sar_paths)+"\n\nClique em EXECUTAR ANÁLISE.")
        self.status.set(f"{len(ps)} produto(s) SAR selecionado(s).")

    def _results(self):
        f=self.tabs[3]; ttk.Label(f,text="Balanço de compartimentos",style="H.TLabel").pack(anchor="w")
        result_box=ttk.Frame(f); result_box.pack(fill="both",expand=True,pady=8)
        self.res=tk.Text(result_box,height=23,wrap="word",yscrollcommand=lambda *a:result_scroll.set(*a))
        result_scroll=ttk.Scrollbar(result_box,orient="vertical",command=self.res.yview)
        self.res.pack(side="left",fill="both",expand=True); result_scroll.pack(side="right",fill="y")
        self._set(self.res,"Clique em EXECUTAR ANÁLISE quando houver dados suficientes.")
    def _sources(self):
        f=self.tabs[4]; ttk.Label(f,text="Rastreabilidade metodológica",style="H.TLabel").pack(anchor="w")
        source_box=ttk.Frame(f); source_box.pack(fill="both",expand=True,pady=8)
        self.src=tk.Text(source_box,height=24,wrap="word",yscrollcommand=lambda *a:source_scroll.set(*a))
        source_scroll=ttk.Scrollbar(source_box,orient="vertical",command=self.src.yview)
        self.src.pack(side="left",fill="both",expand=True); source_scroll.pack(side="right",fill="y")
        txt=("REGRAS DO MOTOR\n• MEDIDO: derivado diretamente do inventário/raster fornecido.\n• MODELADO: proxy/equação publicada, identificado com fonte e domínio.\n• NÃO ESTIMADO: quando não existe suporte defensável.\n\n"
             f"BGB: relação raiz/parte aérea {ROOT_RATIO:.2f}, faixa {ROOT_LOW:.2f}–{ROOT_HIGH:.2f}; {SOURCES['protocol']}.\n"
             f"Conversão biomassa→C: 0,47; {SOURCES['protocol']}.\n"
             f"Solo: {SOURCES['soil']}. O produto PronaSolos utilizado tem resolução nativa de 90 m; o programa preserva essa resolução e não faz falso downscaling.\n"
             f"Necromassa: {SOURCES['deadwood']}; proxy de triagem recebe incerteza elevada e nunca é rotulado como medido.\n"
             "Serrapilheira: proxy só é ativado para Amazônia quando há AGB e é explicitamente rotulado; para MRV recomenda-se amostragem local.")
        self._set(self.src,txt)
    def _set(self,w,t): w.config(state="normal"); w.delete("1.0","end"); w.insert("1.0",t); w.config(state="disabled")
    def _reset_analysis_state(self,keep_geometry=False):
        """Invalida integralmente qualquer resultado derivado da consulta anterior."""
        self._ibge_generation+=1; self._ibge_pending=False; self._pending_execute=False
        old_vector=self.project.get("vector")
        self.project={"version":APP_VERSION}
        if old_vector and keep_geometry:self.project["vector"]=old_vector
        if not keep_geometry:self.gdf=None
        self._map_generation+=1
        self._satellite_map_visible=False
        if getattr(self,"_map_refresh_job",None) is not None:
            try:self.after_cancel(self._map_refresh_job)
            except Exception:pass
            self._map_refresh_job=None
        if hasattr(self,"map_canvas"):
            self.map_canvas.delete("all")
            self.map_canvas.create_text(20,20,anchor="nw",text="Carregue CAR, CCIR ou vetor e visualize o satélite com o polígono.",fill="#455")
        self.soil_raster=None
        self.sar_paths=[]
        self.inv=None
        self.active_source=None
        self.active_input_id=None
        self.biome.set(""); self.phys.set("")
        for widget_name in ("remote_text","res"):
            w=getattr(self,widget_name,None)
            if w is not None:self._set(w,"")
        if hasattr(self,"spatial_text"):self._set(self.spatial_text,"Nova entrada recebida. Resultados anteriores foram descartados.")
        self.status.set("Estado anterior descartado. Preparando nova consulta.")
        if hasattr(self,"vector_status") and not keep_geometry:self.vector_status.set("Nenhum arquivo vetorial carregado.")
        self.update_idletasks()

    def car_lookup(self):
        self._reset_analysis_state()
        try:
            self.status.set("Consultando SICAR..."); self.update_idletasks(); self.gdf=resolve_car(self.car.get()); self.active_source="CAR"; self.active_input_id=self.car.get().strip().upper(); self.ccir.set(""); self._show_geom("SICAR")
        except Exception as e: self.status.set("CAR não resolvido."); messagebox.showwarning("SICAR",str(e))
    def ccir_lookup(self):
        self._reset_analysis_state()
        try:
            self.status.set("Consultando SIGEF pelo código do CCIR..."); self.update_idletasks()
            self.gdf=resolve_ccir_sigef(self.ccir.get()); self.active_source="CCIR"; self.active_input_id=re.sub(r"\D","",self.ccir.get()); self.car.set(""); self._show_geom("CCIR / SIGEF")
        except Exception as e:
            self.status.set("CCIR/SIGEF não resolvido."); messagebox.showwarning("CCIR / SIGEF",str(e))

    def pick_vector(self):
        p=filedialog.askopenfilename(title="Selecionar limite da propriedade",filetypes=[("Vetores","*.kml *.kmz *.geojson *.json *.shp *.gpkg"),("Todos","*.*")])
        if not p:return
        self._reset_analysis_state()
        self.vector_status.set(f"Selecionado: {Path(p).name} — validando arquivo…"); self.status.set("Validando arquivo vetorial..."); self.update_idletasks()
        try:
            self.gdf=read_vector(p); self.project["vector"]=p; self.active_source="VECTOR"; self.active_input_id=str(Path(p).resolve()); self.car.set(""); self.ccir.set("")
            self.vector_status.set(f"Upload concluído ✓  {Path(p).name}  |  {len(self.gdf)} feição(ões) vetorial(is) carregada(s).")
            self.nb.select(self.tabs[1]); self.update_idletasks()
            self._show_geom(Path(p).name); self.vector_status.set(f"Upload concluído ✓  {Path(p).name}  |  {len(self.gdf)} feição(ões) carregada(s).")
        except Exception as e:
            self.gdf=None; self.vector_status.set(f"Upload não concluído — {Path(p).name}: {e}"); self.status.set("Falha ao carregar vetor."); messagebox.showerror("Vetor não carregado",f"O arquivo não foi carregado; nenhuma análise foi iniciada.\n\nArquivo: {Path(p).name}\n\nMotivo: {e}")
    def _schedule_offline_map_fit(self,label):
        """Guarantee a usable map without network: draw now, then redraw after Tk settles."""
        self._satellite_map_visible=False
        if not self._draw_offline_brazil_basemap(label):
            self._draw_aoi_outline(label)
        gen=self._map_generation
        def redraw():
            if self.gdf is None or gen!=self._map_generation:return
            if not self._draw_offline_brazil_basemap(label):
                self._draw_aoi_outline(label)
        self.after_idle(redraw)
        self.after(180,redraw)
        self.after(420,redraw)

    def _show_geom(self,src):
        m=geom_metrics(self.gdf); self.project["geometry_metrics"]=m
        self._set(self.spatial_text,f"Perímetro: {src}\nÁrea geométrica: {m['area_ha']:,.2f} ha\nCentroide: {m['centroid'][1]:.6f}, {m['centroid'][0]:.6f}\nCRS métrico de cálculo: EPSG:{m['utm_epsg']}\n\nPerímetro válido para recorte espacial.")
        self._schedule_offline_map_fit(f"Perímetro carregado: {src} — mapa vetorial offline")
        self.status.set(f"Perímetro carregado: {src}; diagnóstico IBGE em segundo plano."); self.update_idletasks()
        self._ibge_generation+=1; generation=self._ibge_generation; geometry=self.gdf.copy(); self._ibge_pending=True
        def worker():
            try: self._ibge_queue.put((generation,"ok",diagnose_ibge(geometry)))
            except Exception as e: self._ibge_queue.put((generation,"error",str(e)))
        threading.Thread(target=worker,name="EnformIBGE",daemon=True).start()
        self.after(100,self._poll_ibge)
    def _poll_ibge(self):
        try:generation,kind,value=self._ibge_queue.get_nowait()
        except queue.Empty:
            if self.winfo_exists():self.after(120,self._poll_ibge)
            return
        if generation!=self._ibge_generation:return
        self._ibge_pending=False
        if kind!="ok":
            self.project["ibge_diagnosis_error"]=value
            if hasattr(self,"ibge_phys_display"):self.ibge_phys_display.set("Não determinada — falha na consulta ao mapa oficial IBGE: "+str(value)[:160])
            self.status.set("Perímetro carregado; diagnóstico IBGE pendente.")
            self._set(self.spatial_text,self.spatial_text.get("1.0","end").strip()+"\n\nDiagnóstico IBGE pendente: "+value)
            if self._pending_execute:self._pending_execute=False; self.after(0,self.execute)
            return
        d=value; self.project["ibge_diagnosis"]=d
        if d["biomas"]:self.biome.set(d["biomas"][0][0])
        primary=primary_ibge_physiognomy(d)
        if primary:self.phys.set(primary)
        elif d.get("vegetacao_error"):self.phys.set("Não determinada — "+d["vegetacao_error"][:120])
        if hasattr(self,"ibge_biome_display"):
            self.ibge_biome_display.set(self.biome.get() or "Não determinado")
        if hasattr(self,"ibge_phys_display"):
            code=None
            regs=d.get("regioes_fitoecologicas") or []
            if regs and primary and regs[0].get("name")==primary:code=regs[0].get("code")
            shown=self.phys.get() or "Não determinada"
            self.ibge_phys_display.set(shown+(f"  |  código IBGE: {code}" if code else ""))
        btxt="; ".join(f"{n}: {pct:.1f}% ({ha:,.1f} ha)" for n,ha,pct in d["biomas"])
        vtxt=" | ".join(x["campo"]+": "+"; ".join(f"{n}: {pct:.1f}% ({ha:,.1f} ha)" for n,ha,pct in x["classes"][:8]) for x in d["vegetacao"]) or ("PENDENTE: "+d.get("vegetacao_error","sem classe"))
        code_txt="; ".join(f"{x['code'] or 'código N/D'} — {x['name']}: {x['percent']:.1f}% ({x['area_ha']:,.1f} ha)" for x in d.get("regioes_fitoecologicas",[])[:8])
        if code_txt:vtxt="Região fitoecológica (IBGE legenda_1): "+code_txt+" | "+vtxt
        self._set(self.spatial_text,self.spatial_text.get("1.0","end").strip()+"\n\nIBGE — Bioma(s): "+btxt+"\nIBGE 2026 — Vegetação: "+vtxt)
        self.status.set("Perímetro e diagnóstico IBGE concluídos." if not d.get("vegetacao_error") else "Bioma IBGE concluído; fitofisionomia não classificada — modelos que exigem classe IBGE ficam bloqueados; apenas rotas independentes da classe podem prosseguir.")
        if self._pending_execute:self._pending_execute=False; self.after(0,self.execute)
    def pick_soil(self):
        p=filedialog.askopenfilename(filetypes=[("GeoTIFF","*.tif *.tiff")])
        if p:self.soil_raster=p; self.status.set("Raster de COS selecionado.")
    def auto_soil(self):
        if self.gdf is None:return messagebox.showwarning("Solo","Carregue/resolva o perímetro primeiro.")
        try:self.status.set("Baixando COS Embrapa..."); self.update_idletasks(); self.soil_raster=try_download_embrapa_soc(self.gdf); self.status.set("COS Embrapa obtido.")
        except Exception as e:self.status.set("COS automático indisponível."); messagebox.showwarning("Solo Embrapa",str(e))
    def execute(self,event=None):
        if self._analysis_running:
            self.status.set("Análise já em execução; aguarde.")
            return
        if self._ibge_pending:
            self._pending_execute=True; self.status.set("A AOI está carregada. A análise começará assim que o diagnóstico territorial do IBGE terminar.")
            return
        car=self.car.get().strip().upper(); ccir=re.sub(r"\D","",self.ccir.get())
        if car and ccir:return messagebox.showwarning("Identificação","Informe CAR ou CCIR, não ambos. Para outro perímetro, carregue um arquivo vetorial.")
        if car and (self.active_source!="CAR" or self.active_input_id!=car):
            self._reset_analysis_state(); self.status.set("Consultando automaticamente o CAR no SICAR…"); self.update_idletasks()
            try:
                self.gdf=resolve_car(car); self.active_source="CAR"; self.active_input_id=car; self.ccir.set(""); self._show_geom("SICAR — consulta automática")
            except Exception as e:
                self.status.set("Consulta automática CAR/SICAR não concluída."); return messagebox.showwarning("SICAR",str(e))
        elif ccir and (self.active_source!="CCIR" or self.active_input_id!=ccir):
            self._reset_analysis_state(); self.status.set("Consultando automaticamente CCIR/SNCR e perímetro SIGEF…"); self.update_idletasks()
            try:
                self.gdf=resolve_ccir_sigef(ccir); self.active_source="CCIR"; self.active_input_id=ccir; self.car.set(""); self._show_geom("CCIR/SNCR — SIGEF automático")
            except Exception as e:
                self.status.set("Consulta automática CCIR/SNCR/SIGEF não concluída."); return messagebox.showwarning("CCIR / SIGEF",str(e))
        if self._ibge_pending:
            self._pending_execute=True; self.status.set("Perímetro resolvido. A análise começará após o diagnóstico territorial do IBGE.")
            return
        if self.gdf is None:return messagebox.showwarning("Perímetro necessário","Informe CAR, CCIR ou carregue um arquivo vetorial; depois pressione EXECUTAR ANÁLISE.")
        if self.sar_paths:return self._execute_main(event)
        # Snapshot every Tk variable on the GUI thread before starting the worker.
        gdf=self.gdf.copy(); biome=self.biome.get(); phys=self.phys.get(); token=self.esa_token.get().strip()
        edl_token=self.edl_token.get().strip(); edl_user=self.edl_user.get().strip(); edl_password=self.edl_password.get(); cdse_token=self.cdse_token.get().strip()
        cdse_client_id=self.cdse_client_id.get().strip(); cdse_client_secret=self.cdse_client_secret.get().strip()
        self._analysis_running=True; self.pipeline_btn.state(["disabled"]); self.global_execute_btn.state(["disabled"])
        self.status.set("Consultando e processando SAR em segundo plano…")
        def worker():
            try:
                sar=automatic_pipeline(gdf,biome,phys,token,edl_user=edl_user,edl_password=edl_password,edl_token=edl_token,cdse_token=cdse_token,cdse_client_id=cdse_client_id,cdse_client_secret=cdse_client_secret)
                try: soil_profiles=pronasolos_soc_profiles(gdf)
                except Exception as e: soil_profiles={"error":str(e)}
                self._analysis_queue.put(("ok",{"sar":sar,"soil":soil_profiles}))
            except Exception:self._analysis_queue.put(("error",traceback.format_exc()))
        threading.Thread(target=worker,name="EnformAnalysis",daemon=True).start()
        self.after(120,self._poll_analysis)

    def _poll_analysis(self):
        try:kind,payload=self._analysis_queue.get_nowait()
        except queue.Empty:
            if self._analysis_running:self.after(120,self._poll_analysis)
            return
        self._analysis_running=False; self.pipeline_btn.state(["!disabled"]); self.global_execute_btn.state(["!disabled"])
        if kind=="error":
            log=Path.home()/".enform_verde"/"enform_diagnostico.log"; log.parent.mkdir(parents=True,exist_ok=True); log.write_text(payload,encoding="utf-8")
            self.status.set("Falha controlada — programa permanece responsivo.")
            return messagebox.showerror("Análise","Falha controlada. Log gravado em:\n"+str(log))
        self.status.set("SAR consultado; calculando carbono…")
        self.after(1,lambda:self._execute_main(precomputed_sar=payload.get("sar"),precomputed_soil=payload.get("soil")))

    def _execute_main(self,event=None,precomputed_sar=None,precomputed_soil=None):
        # Cada execução substitui, nunca acumula, os resultados derivados da geometria corrente.
        for k in ("analysis_rows","area_ha","total_tc_ha","total_tco2_ha","last_result"):
            self.project.pop(k,None)
        current_car=self.car.get().strip().upper()
        current_ccir=re.sub(r"\D","",self.ccir.get())
        if current_car and (self.active_source!="CAR" or self.active_input_id!=current_car):
            self._reset_analysis_state()
            try:
                self.gdf=resolve_car(current_car); self.active_source="CAR"; self.active_input_id=current_car; self._show_geom("SICAR")
            except Exception as e:return messagebox.showwarning("Perímetro necessário",str(e))
        elif current_ccir and (self.active_source!="CCIR" or self.active_input_id!=current_ccir):
            self._reset_analysis_state()
            try:
                self.gdf=resolve_ccir_sigef(current_ccir); self.active_source="CCIR"; self.active_input_id=current_ccir; self._show_geom("CCIR / SIGEF")
            except Exception as e:return messagebox.showwarning("Perímetro necessário",str(e))
        if self.gdf is None:return messagebox.showwarning("Perímetro necessário","Informe CAR/CCIR ou carregue o arquivo vetorial da propriedade.")
        try:
            self.status.set("Executando estimativa remota..."); self.update_idletasks()
            area=geom_metrics(self.gdf)["area_ha"]
            if precomputed_sar is not None:
                sar=precomputed_sar
            elif self.sar_paths:
                sar=process_real_sar(self.gdf,self.sar_paths,self.biome.get(),self.phys.get())
            else:
                raise RuntimeError("Pipeline automático sem resultado do worker.")
            self.project["sar_result"]=sar
            # Contract: a numerical AGB is a SAR result only when provenance explicitly says SAR.
            if sar.get("agb_mg_ha") is not None and not str(sar.get("data_origin","")).startswith("SAR"):
                self.project["reference_result"]=sar
                raise RuntimeError("Contrato de proveniência violado: AGB numérica sem origem SAR. O valor foi bloqueado para impedir rotulagem incorreta.")
            if sar.get("agb_mg_ha") is None:
                msg=sar.get("message") or sar.get("reason") or "AGB não pôde ser estimada."
                lit=sar.get("literature_reference") or {}
                if lit.get("available"):
                    # Always deliver an analysis, but never relabel literature as SAR.
                    sar=dict(lit); sar["data_origin"]=lit.get("data_origin") or ("LITERATURA_MICRORREGIONAL" if lit.get("data_origin")=="LITERATURA_MICRORREGIONAL" else "LITERATURA_SECUNDARIA")
                    if sar["data_origin"]=="MODELAGEM_LITERATURA_HIERARQUICA":
                        sar["status"]=("FALLBACK MODELADO HIERÁRQUICO — SAR PROCESSADO, SEM MODELO AGB" if lit.get("sar_processed") else "FALLBACK MODELADO HIERÁRQUICO — SAR NÃO PROCESSADO")
                    else:
                        sar["status"]=("FALLBACK MICRORREGIONAL — SAR PROCESSADO, SEM MODELO AGB" if lit.get("sar_processed") else "FALLBACK DE REFERÊNCIA — SAR NÃO PROCESSADO")
                    sar["source"]=lit.get("source","biblioteca científica interna"); sar["sar_diagnostic"]=msg
                    sar["uncertainty_mg_ha"]=float(lit.get("uncertainty_mg_ha",float(lit.get("agb_mg_ha",0))*float(lit.get("uncertainty_pct",30))/100))
                    sar["uncertainty_kind"]=lit.get("uncertainty_kind", "amplitude bibliográfica; não IC95%")
                    sar["audit"]=((self.project.get("sar_result") or {}).get("audit") or {})
                    sar["sar_validation_metrics"]={"RMSE":"N/D","MAE":"N/D","viés":"N/D","R²":"N/D","motivo":"não há predições AGB SAR pareadas com parcelas independentes compatíveis"}
                    self.project["sar_result"]=sar
                    self.project["sar_warning"]=msg
                else:
                    # Report actual SAR operations while withholding unsupported AGB/carbon numbers.
                    audit=sar.get("audit") or {}; lines=["RELATÓRIO SAR — DADOS PROCESSADOS; AGB NÃO ESTIMADA",f"Projeto: {self.name.get()}",f"Área da AOI: {area:,.2f} ha",f"Bioma: {self.biome.get() or 'não determinado'} | Fitofisionomia: {self.phys.get() or 'não determinada'}","", "Motivo: "+str(msg), "", "O SAR foi processado, mas o catálogo não contém equação validada compatível com os preditores e o domínio desta AOI. Não se publica AGB nem carbono sem suporte defensável.","", "PRODUTOS E PIXELS PROCESSADOS:"]
                    processed=audit.get("processed_without_agb") or []
                    for item in processed:
                        lines.append(f"• {item.get('source','SAR')} | {item.get('provider','provedor não informado')}")
                        for z in item.get("stats") or []: lines.append("  "+json.dumps(z,ensure_ascii=False,sort_keys=True))
                        if item.get("scene_ids"): lines.append("  Cenas: "+", ".join(map(str,item["scene_ids"])))
                    for z in sar.get("stats") or []: lines.append("• raster processado: "+json.dumps(z,ensure_ascii=False,sort_keys=True))
                    for w in audit.get("warnings") or []: lines.append("Aviso: "+str(w))
                    lines += ["", "Métricas de validação AGB: RMSE, MAE, viés e R² não são aplicáveis sem modelo treinado/validado compatível. Incerteza da AGB: não estimável.", "Estado: PROCESSAMENTO SAR REAL CONCLUÍDO; ESTIMATIVA AGB PENDENTE DE MODELO/PARCELAS COMPATÍVEIS."]
                    report="\n".join(lines); self.project["partial_sar_analysis"]={"area_ha":area,"sar_result":sar,"report":report}; self.project["area_ha"]=area; self.project["last_result"]=report
                    self._set(self.remote_text,"SAR processado na AOI. AGB não estimada por falta de modelo validado compatível; consulte a trilha, os pixels processados e o motivo no relatório.")
                    self._set(self.res,report); self.nb.select(self.tabs[3]); self.status.set("Processamento SAR concluído; AGB não estimada sem calibração compatível."); return
            agb=float(sar["agb_mg_ha"]); sar_unc=float(sar.get("uncertainty_mg_ha") or 0.0)
            unc_kind=str(sar.get("uncertainty_kind") or "incerteza do produto/modelo")
            unc_mult=1.0 if ("amplitude bibliográfica" in unc_kind or "intervalo preditivo aproximado" in unc_kind) else 1.96
            explicit_range=sar.get("agb_range_mg_ha")
            if isinstance(explicit_range,(list,tuple)) and len(explicit_range)==2:
                agb_lo,agb_hi=map(float,explicit_range)
            else:agb_lo=max(0.0,agb-unc_mult*sar_unc); agb_hi=agb+unc_mult*sar_unc
            agc=agb*CARBON_FRACTION
            agc_lo=agb_lo*CARBON_FRACTION; agc_hi=agb_hi*CARBON_FRACTION
            soil_profiles=precomputed_soil if isinstance(precomputed_soil,dict) else {}
            soil_error=soil_profiles.get("error") if soil_profiles else "PronaSolos não retornou perfil."
            parts=[
              ("Biomassa aérea",agc,sar.get("status","SAR PROCESSADO"),f"AGB={agb:,.1f} Mg/ha; incerteza={sar_unc:,.1f} Mg/ha; carbono={CARBON_FRACTION:.2f}; faixa C={agc_lo:,.2f}–{agc_hi:,.2f} tC/ha","produto SAR ou referência secundária, conforme origem",sar.get("source",sar.get("status","produto processado")))]
            regional_components=sar.get("regional_components") or {}
            for cname,component in regional_components.items():
                dry=float(component["mean_dry_mg_ha"]); bounds=list(map(float,component["range_dry_mg_ha"]))
                note=(f"biomassa seca={dry:,.2f} Mg/ha; faixa descritiva={bounds[0]:,.2f}–{bounds[1]:,.2f} Mg/ha; "
                      f"fração C operacional={CARBON_FRACTION:.2f}; a faixa não é IC95% nem erro SAR. "+str(component.get("method","")))
                parts.append((cname,dry*CARBON_FRACTION,"REFERÊNCIA MICRORREGIONAL",note,
                              "estoque de referência publicado; não é predição SAR nem medição da AOI",component.get("source","literatura científica regional")))
            p030=soil_profiles.get("0–30 cm") if soil_profiles else None
            if p030:
                parts.append(("Solo 0–30 cm",p030["tc_ha"],"MAPEAMENTO DIGITAL",f'{p030["n_samples"]} amostras do mapa 90 m; DP espacial {p030["spatial_sd_tc_ha"]:,.2f} tC/ha',"PronaSolos 90 m: soma 0–5 + 5–15 + 15–30 cm","Embrapa Solos/PronaSolos"))
            total=sum(x[1] for x in parts); co2=total*44/12
            # Deeper SOC profiles are reported independently and are NOT summed again into Carbono Total.
            for depth in ("0–60 cm","0–100 cm","0–200 cm"):
                p=soil_profiles.get(depth) if soil_profiles else None
                if p:parts.append((f"Solo {depth}",p["tc_ha"],"MAPEAMENTO DIGITAL",f'{p["n_samples"]} amostras do mapa 90 m; DP espacial {p["spatial_sd_tc_ha"]:,.2f} tC/ha; não somado novamente ao Carbono Total',f"PronaSolos 90 m: soma das camadas até {depth.split('–')[1]}","Embrapa Solos/PronaSolos"))
            if soil_error:self.project["soil_warning"]=soil_error
            # Statistical/uncertainty metadata. Never label a descriptive range as a confidence interval.
            agb_abs=(sar_unc*CARBON_FRACTION) if sar_unc else None
            agb_pct=(sar_unc/agb*100) if agb and sar_unc else None
            agb_metric=sar.get("uncertainty_kind","incerteza do produto/modelo")
            rows=[]
            for name,val,status,note,method,source in parts:
                origem=((str(sar.get("data_origin") or "NÃO CLASSIFICADO") if name=="Biomassa aérea" else ("MAPEAMENTO" if name.startswith("Solo ") else ("LITERATURA_MICRORREGIONAL" if name in regional_components else "MODELADO"))))
                if name=="Biomassa aérea":
                    ea,ep,metric,level=agb_abs,agb_pct,agb_metric,(("envelope descritivo; sem cobertura probabilística declarada" if "envelope descritivo" in agb_metric else ("faixa bibliográfica; não IC95%" if "bibliográfica" in agb_metric else "1σ/DP ou métrica do produto/modelo")) if sar_unc else "N/D")
                elif name in regional_components:
                    c=regional_components[name]; bounds=list(map(float,c["range_dry_mg_ha"]))
                    ea=max(float(c["mean_dry_mg_ha"])-bounds[0],bounds[1]-float(c["mean_dry_mg_ha"]))*CARBON_FRACTION
                    ep=(ea/val*100 if val else None)
                    metric="envelope descritivo da fonte × fração C operacional; não é IC95%, erro preditivo ou erro SAR"
                    level="dispersão/faixa de estudos locais; sem cobertura probabilística declarada"
                elif name=="Biomassa subterrânea":
                    ea=None; ep=None; metric="N/D — sem dados compatíveis"; level="não estimado"
                elif name.startswith("Solo "):
                    depth=name.replace("Solo ",""); sp=soil_profiles.get(depth,{})
                    sd=float(sp.get("spatial_sd_tc_ha",0.0)); ea=sd; ep=(sd/val*100 if val else None)
                    metric=f"DP espacial={sd:,.2f} tC/ha ({ep:.1f}% da média)" if ep is not None else "DP espacial N/D"
                    level="variabilidade espacial do mapa; não IC95% nem erro de predição"
                else:
                    ea=None; ep=None; metric="proxy bibliográfico/modelado sem distribuição de erro validada"; level="erro estatístico N/D"
                rows.append({"parametro":name,"tc":val,"tco2":val*44/12,"origem":origem,"status":status,"metodo":method,"fonte":source,"obs":note,
                             "erro_abs_tc":ea,"erro_pct":ep,"erro_metrica":metric,"nivel_confianca":level})
            # Propagate only quantified independent 1-sigma components; report coverage of uncertainty.
            q=[r for r in rows if r.get("erro_abs_tc") is not None and r.get("origem") not in ("LITERATURA_MICRORREGIONAL","LITERATURA_SECUNDARIA","MODELAGEM_LITERATURA_HIERARQUICA")]
            total_sigma=math.sqrt(sum(r["erro_abs_tc"]**2 for r in q)) if q else None
            total_err_pct=(total_sigma/total*100) if total_sigma is not None and total else None
            self.project["total_uncertainty"]={"sigma_tc_ha":total_sigma,"pct":total_err_pct,"quantified_components":len(q),"total_components":len(rows),
                "note":"propagação RSS dos componentes quantificados; não inclui componentes com erro estatístico N/D"}
            self.project["analysis_rows"]=rows
            self.project["area_ha"]=area; self.project["total_tc_ha"]=total; self.project["total_tco2_ha"]=co2
            audit=sar.get("audit") or {}
            diag=[]
            if audit:
                diag=["","TRILHA SAR:"]
                bb=audit.get("biomass_l2b") or {}; diag.append(f"BIOMASS P L2B catalogado: {bb.get(\'count\',0)} | operacional={bb.get(\'operational_count\',0)} | IOC={bb.get(\'ioc_count\',0)}")
                cc=audit.get("cci") or {}; diag.append(f"CCI AGB: {cc.get('downloaded',0)} arquivo(s) baixado(s)" if isinstance(cc,dict) else "CCI AGB: não disponível")
                ad=audit.get("asf_download") or {}
                if ad: diag.append(f"ASF/NISAR/ALOS: cena={ad.get('scene')} | pré-processamento={ad.get('preprocess')} | candidatos={ad.get('candidate_count')}")
                matrix=audit.get("national_route_matrix") or {}
                if matrix:
                    diag.append("COBERTURA PREDITIVA NACIONAL: "+str(matrix.get("local_numeric_state")))
                    diag.append("Escopo implementado: "+", ".join(matrix.get("supported_scope") or []))
                    diag.append("Política: "+str(matrix.get("policy")))
                    ifn=matrix.get("ifn_sfb_reference") or {}
                    if ifn:
                        diag.append("IFN/SFB: referência nacional/estadual para aferição e alometria; não é raster local da AOI. Fonte: "+str(ifn.get("url")))
                for item in audit.get("processed_without_agb",[]):
                    diag.append(f"Pixel SAR PROCESSADO sem equação AGB compatível: {item.get('source','SAR')} | {item.get('provider','provedor não informado')}")
                    for z in item.get("stats",[]): diag.append("  "+json.dumps(z,ensure_ascii=False,sort_keys=True))
                blocker=audit.get("agb_blocker") or sar.get("agb_blocker") or {}
                if blocker:
                    diag.append("BLOQUEIO AGB-SAR: "+str(blocker.get("primary_blocker")))
                    cm=blocker.get("closest_model") or {}
                    if cm:
                        diag.append("  modelo mais próximo: "+str(cm.get("model_id"))+" | sensor="+str(cm.get("sensor")))
                        avail=cm.get("available_predictors") or []
                        miss=cm.get("missing_predictors") or []
                        diag.append("  preditores disponíveis: "+(", ".join(avail) if avail else "nenhum dos exigidos"))
                        diag.append("  preditores faltantes: "+(", ".join(miss) if miss else "nenhum"))
                        diag.append("  restrição: "+str(cm.get("constraints")))
                ldiag=audit.get("lband_dualpol_diagnostic") or {}
                if ldiag:
                    diag += [
                        "DIAGNÓSTICO L-BAND DUAL-POL:",
                        f"  papel={ldiag.get('role')} | saturação={ldiag.get('saturation_risk')} | AGB quantitativa permitida={ldiag.get('quantitative_agb_from_dualpol_permitted')}",
                        f"  HH={float(ldiag.get('hh_gamma0_db',float('nan'))):.2f} dB | HV={float(ldiag.get('hv_gamma0_db',float('nan'))):.2f} dB | HH-HV={float(ldiag.get('hh_minus_hv_db',float('nan'))):.2f} dB | HV/HH={float(ldiag.get('hv_over_hh_linear',float('nan'))):.3f} | RFDI={float(ldiag.get('rfdi',float('nan'))):.3f}",
                        "  decisão: "+str(ldiag.get("reason")),
                    ]
                    for ref in ldiag.get("references",[]):
                        diag.append("  referência: "+str(ref.get("source"))+" | DOI "+str(ref.get("doi"))+" | "+str(ref.get("note")))
                for w in audit.get("warnings",[]): diag.append("Aviso: "+str(w))
            if sar.get("data_origin") in ("LITERATURA_MICRORREGIONAL","MODELAGEM_LITERATURA_HIERARQUICA"):
                diag += ["", "MÉTRICAS DE VALIDAÇÃO SAR: RMSE=N/D; MAE=N/D; viés=N/D; R²=N/D — faltam pares independentes parcela–pixel SAR.",
                         "A estimativa regional é um resumo publicado e não gera raster/mapa AGB pixel a pixel.",
                         "Incerteza: "+str(sar.get("uncertainty_kind")),
                         "Suporte: "+str(sar.get("n_plots"))+" parcelas resumidas em "+str(sar.get("n_independent_sites"))+" sítios; distância ao km 83 = "+f"{float(sar.get('distance_from_km83_km',float('nan'))):.2f} km."]
            lines=([f"SAR PROCESSADO — AGB NÃO DERIVADA DO SAR: {self.project.get('sar_warning')}",""] if self.project.get("sar_warning") else [])+[f"ENFORM VERDE {APP_VERSION}",f"Projeto: {self.name.get()}",f"Sensor/produto: {self.sensor.get()}",f"Bioma IBGE: {self.biome.get()} | Fitofisionomia IBGE (legenda_1): {self.phys.get()}",f"Área analisada: {area:,.2f} ha",""]+diag
            for r in rows:
                err=(f"±{r['erro_abs_tc']:.2f} tC/ha ({r['erro_pct']:.1f}%)" if r.get('erro_pct') is not None else "N/D")
                lines += [f"{r['parametro']}",f"  {r['tc']:,.2f} tC/ha  |  {r['tco2']:,.2f} tCO₂e/ha",f"  ORIGEM DO DADO: {r['origem']}",f"  Erro/incerteza: {err}",f"  Nível estatístico: {r['nivel_confianca']}",f"  Métrica: {r['erro_metrica']}",f"  Método/produto: {r['metodo']}",f"  Fonte: {r['fonte']}",f"  {r['status']} — {r['obs']}",""]
            if not p030: lines += ["Solo 0–30 cm","  NÃO CALCULADO — PronaSolos não retornou as três camadas necessárias nesta execução.","  Diagnóstico: "+str(soil_error),""]
            missing=[]
            if "Biomassa subterrânea" not in regional_components: missing.append("biomassa subterrânea")
            if not any(n.startswith("Necromassa") for n in regional_components): missing.append("necromassa")
            if not any(n.startswith("Serapilheira") for n in regional_components): missing.append("serapilheira")
            lines += ["COMPARTIMENTOS NÃO SOMADOS", ("Nenhum compartimento adicional elegível ficou sem estimativa nesta execução." if not missing else "; ".join(missing)+": não estimada por falta de dados/modelos regionais compatíveis."),
                      *( ["Referências microrregionais são benchmarks secundários; não equivalem a medição da AOI nem a modelos alométricos/SAR calibrados."] if regional_components else [] ), "",
                      "TOTAL DOS COMPARTIMENTOS DISPONÍVEIS",f"  {total:,.2f} tC/ha  |  {co2:,.2f} tCO₂e/ha",f"  Total na área: {total*area:,.0f} tC  |  {co2*area:,.0f} tCO₂e","",
                      "QUALIDADE: resultado de triagem/planejamento remoto. O relatório distingue SAR efetivamente processado de referência bibliográfica secundária e não substitui inventário de campo."]
            self._set(self.remote_text,f"Biomassa aérea: {agc:,.2f} tC/ha | {agc*44/12:,.2f} tCO₂e/ha\nFaixa de referência: {agc_lo:,.2f}–{agc_hi:,.2f} tC/ha | {agc_lo*44/12:,.2f}–{agc_hi*44/12:,.2f} tCO₂e/ha")
            self._set(self.res,"\n".join(lines)); self.project["last_result"]="\n".join(lines); self.nb.select(self.tabs[3]); self.status.set("Estimativa remota concluída.")
        except Exception as e:self.status.set("Falha."); messagebox.showerror("Análise",str(e))

    def export_excel(self):
        if not self.project.get("analysis_rows"):
            return messagebox.showwarning("Excel","Execute a análise antes de exportar.")
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",filetypes=[("Excel","*.xlsx")])
        if not p:return
        wb=Workbook(); orange="F29A00"; green="0B3D2E"; white="FFFFFF"; pale="F4F6F5"; line="D5DDD9"
        thin=Side(style="thin",color=line)
        headers=["Categoria","Parâmetro / resultado","tC/ha","tCO₂e/ha","Status","Método","Fonte","Observação"]
        def setup(sh,title):
            sh.sheet_view.showGridLines=False; sh.freeze_panes="A4"; sh.auto_filter.ref="A3:H200"
            sh.merge_cells("A1:H1"); c=sh["A1"]; c.value=title; c.font=Font(size=18,bold=True,color=white); c.fill=PatternFill("solid",fgColor=green); c.alignment=Alignment(vertical="center")
            sh.row_dimensions[1].height=32
            for col,w in zip("ABCDEFGH",[20,34,16,18,25,28,30,55]): sh.column_dimensions[col].width=w
            for j,h in enumerate(headers,1):
                c=sh.cell(3,j,h); c.font=Font(bold=True,color=white); c.fill=PatternFill("solid",fgColor=orange); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
            sh.row_dimensions[3].height=28
        def put(sh,data):
            for i,row in enumerate(data,4):
                sh.row_dimensions[i].height=32
                for j,v in enumerate(row,1):
                    c=sh.cell(i,j,v); c.fill=PatternFill("solid",fgColor=(white if i%2==0 else pale)); c.border=Border(top=thin,bottom=thin,left=thin,right=thin); c.alignment=Alignment(vertical="center",wrap_text=True)
                    if j in (3,4) and isinstance(v,(int,float)): c.number_format='#,##0.00'
                sh.cell(i,2).font=Font(bold=True,color=green)
        area=self.project["area_ha"]; total=self.project["total_tc_ha"]; totalco2=self.project["total_tco2_ha"]; ar=self.project["analysis_rows"]
        ws=wb.active; ws.title="Resumo Executivo"; setup(ws,"Enform Verde — Resumo Executivo")
        put(ws,[["Entrada","Projeto",None,None,"Informado","cadastro",None,self.name.get()],
                ["Diagnóstico IBGE","Bioma dominante",None,None,"Calculado espacialmente","interseção de polígonos","IBGE — Biomas 2025",self.biome.get()],
                ["Diagnóstico IBGE","Fitofisionomia/região fitoecológica dominante",None,None,"Calculado espacialmente","interseção de polígonos","IBGE — Vegetação 2026",self.phys.get()],
                ["Entrada","Área analisada",None,None,"Calculado","geometria","CAR/vetor",f"{area:,.2f} ha"],
                ["Entrada","Sensor/produto selecionado",None,None,"Informado","SAR/multissensor","ESA/fornecedor",self.sensor.get()],
                ["Resultado","Carbono total por hectare",total,totalco2,"CONSOLIDADO","soma dos compartimentos","Enform","somente compartimentos disponíveis"],
                ["Resultado","Carbono total da propriedade",None,None,"CONSOLIDADO","total/ha × área","Enform",f"{total*area:,.0f} tC | {totalco2*area:,.0f} tCO₂e"]])
        sh=wb.create_sheet("Compartimentos"); setup(sh,"Enform Verde — Compartimentos de carbono")
        put(sh,[["Resultado — "+r.get("origem","N/D"),r["parametro"],r["tc"],r["tco2"],r["status"],r["metodo"],r["fonte"],r["obs"]] for r in ar])
        for sheet,param in [("Biomassa Aérea","Biomassa aérea"),("Biomassa Subterrânea","Biomassa subterrânea"),("Necromassa","Necromassa"),("Serapilheira","Serapilheira")]:
            sh=wb.create_sheet(sheet); setup(sh,"Enform Verde — "+sheet)
            rr=[r for r in ar if r["parametro"]==param]
            data=[["Resultado",r["parametro"],r["tc"],r["tco2"],r["status"],r["metodo"],r["fonte"],r["obs"]] for r in rr]
            if not data:data=[["Resultado",param,None,None,"NÃO CALCULADO","—","—","Não houve dado válido nesta execução; nenhum valor foi inventado."]]
            put(sh,data)
        sh=wb.create_sheet("Carbono do Solo"); setup(sh,"Enform Verde — Carbono orgânico do solo")
        soil_rows=[r for r in ar if r["parametro"].startswith("Solo ")]
        soil_data=[["MAPEAMENTO — PronaSolos",r["parametro"],r["tc"],r["tco2"],r["status"],r["metodo"],r["fonte"],r["obs"]+" | "+r.get("erro_metrica","")] for r in soil_rows]
        if not soil_data: soil_data=[["MAPEAMENTO","Carbono do solo",None,None,"NÃO CALCULADO","PronaSolos 90 m","Embrapa Solos/PronaSolos","Sem valores válidos nesta execução."]]
        put(sh,soil_data)
        sh=wb.create_sheet("Diagnóstico IBGE"); setup(sh,"Enform Verde — Diagnóstico territorial IBGE")
        diag=self.project.get("ibge_diagnosis",{})
        idata=[]
        for n,ha,pct in diag.get("biomas",[]): idata.append(["Bioma",n,None,None,"IBGE oficial","interseção espacial","IBGE — Biomas 2025",f"{ha:,.2f} ha | {pct:.2f}% da área"])
        for grp in diag.get("vegetacao",[]):
            for n,ha,pct in grp.get("classes",[]): idata.append(["Vegetação "+grp.get("campo",""),n,None,None,"IBGE oficial 2026","interseção espacial","IBGE — Vegetação 2026",f"{ha:,.2f} ha | {pct:.2f}% da área"])
        if not idata:idata=[["Diagnóstico","IBGE",None,None,"NÃO DISPONÍVEL","—","IBGE","A consulta/interseção não foi concluída nesta execução."]]
        put(sh,idata)
        sh=wb.create_sheet("Sensores SAR"); setup(sh,"Enform Verde — Sensores SAR")
        put(sh,[["Sensor","ESA Biomass — banda P",None,None,"Preferencial","PolSAR/PolInSAR/TomoSAR","ESA","Primeiro SAR orbital em banda P; usar somente quando produto efetivamente processado."],
                ["Sensor","ALOS/PALSAR / ALOS-2",None,None,"Histórico/complementar","banda L","JAXA","Séries históricas para atributos estruturais."],
                ["Sensor","TerraSAR-X / TanDEM-X / COSMO-SkyMed",None,None,"Complementar","banda X","Operadores","Textura e estrutura do dossel; não é banda do satélite Biomass."]])
        sh=wb.create_sheet("Modelos e QA"); setup(sh,"Enform Verde — Modelos, QA e incerteza")
        put(sh,[["QA","Mensuração SAR",None,None,"REGRA","controle metodológico","Enform","Não declarar mensuração SAR sem produto efetivamente processado."],
                ["Modelo","Conversão C→CO₂e",None,None,"Aplicado","tC × 44/12","estequiometria","Todas as estimativas de carbono são exibidas em tC/ha e tCO₂e/ha."],
                ["Modelo","Biblioteca brasileira",None,None,"Prioritária","seleção por domínio","IFN/SFB + Embrapa","Bioma, fitofisionomia e domínio de calibração devem ser compatíveis."]])
        sh=wb.create_sheet("Bibliografia"); setup(sh,"Enform Verde — Bibliografia e proveniência")
        put(sh,[["Fonte","ESA Biomass",None,None,"Oficial","P-band SAR","ESA","Missão orbital para biomassa florestal."],
                ["Fonte","IFN / Painel de Biomassa e Carbono",None,None,"Prioritária","inventário/equações","SFB","Base brasileira para seleção e validação."],
                ["Fonte","Embrapa",None,None,"Prioritária","protocolos/equações/COS","Embrapa","Fontes brasileiras priorizadas no motor."]])
        wb.save(p); self.status.set("Excel exportado com sucesso.")

    def save_report(self):
        txt=self.res.get("1.0","end").strip()
        if not txt:return
        p=filedialog.asksaveasfilename(defaultextension=".txt",filetypes=[("Relatório TXT","*.txt")])
        if p:Path(p).write_text(txt+"\n\nFONTES\n"+self.src.get("1.0","end"),encoding="utf-8"); self.status.set("Relatório salvo.")

if __name__=="__main__":
    if "--acceptance-test" in sys.argv:
        try:acceptance_test(sys.argv[sys.argv.index("--acceptance-test")+1])
        except Exception:
            (Path(tempfile.gettempdir())/"enform_verde_acceptance_error.txt").write_text(traceback.format_exc(),encoding="utf-8")
            sys.exit(1)
    elif "--ui-smoke" in sys.argv:ui_smoke_test()
    elif "--self-test" in sys.argv:self_test()
    else:App().mainloop()
